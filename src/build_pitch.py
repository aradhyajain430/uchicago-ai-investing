"""SPROUTs-style 10-slide LONG LITE pitch for the UChicago competition."""
from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from lxml import etree
from copy import deepcopy

ROOT = Path(r"C:\Users\Dell\Desktop\uchicagoaicomp")
OUT = ROOT / "final_slideshow" / "LITE_Round1_Pitch.pptx"
CHARTS = ROOT / "final_slideshow" / "charts"

# Sprouts palette
G = RGBColor(0x2F, 0x7A, 0x38)
G2 = RGBColor(0x3D, 0x8B, 0x40)
GL = RGBColor(0xC8, 0xE6, 0xC9)
GP = RGBColor(0xE8, 0xF5, 0xE9)
GD = RGBColor(0x1B, 0x4D, 0x1E)
WH = RGBColor(0xFF, 0xFF, 0xFF)
BK = RGBColor(0x22, 0x22, 0x22)
GY = RGBColor(0x55, 0x55, 0x55)
GO = RGBColor(0xC9, 0xA2, 0x27)
RD = RGBColor(0xC6, 0x28, 0x28)
BE = RGBColor(0xF5, 0xF5, 0xF5)

W, H = Inches(13.333), Inches(7.5)


def rgb_hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


def set_run(run, text, size=11, bold=False, color=BK, font="Calibri", italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font


def add_tb(slide, l, t, w, h, text, size=11, bold=False, color=BK, align=PP_ALIGN.LEFT,
           font="Calibri", italic=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    set_run(p.add_run() if p.runs else _ensure_run(p), text, size, bold, color, font, italic)
    if not p.runs:
        set_run(p.add_run(), text, size, bold, color, font, italic)
    else:
        # first run may be empty default
        if p.runs[0].text == "" and len(p.runs) > 1:
            pass
    return box


def _ensure_run(p):
    if p.runs:
        return p.runs[0]
    return p.add_run()


def tb(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, bold, color, align, italic) or str."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        p.space_before = Pt(0)
        if isinstance(line, str):
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            set_run(r, line, 11, False, BK)
        else:
            text, size, bold, color = line[:4]
            align = line[4] if len(line) > 4 else PP_ALIGN.LEFT
            italic = line[5] if len(line) > 5 else False
            p.alignment = align
            r = p.add_run()
            set_run(r, text, size, bold, color, italic=italic)
    return box


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def footer(slide, n, total=10):
    rect(slide, Inches(0), Inches(7.28), W, Inches(0.22), G)
    tb(slide, Inches(0.28), Inches(7.28), Inches(11.5), Inches(0.22), [
        ("UChicago AI-Enabled Investing Competition  |  18 September 2026  |  LONG LITE  |  Horizon 3–12 months  |  Sources on slide 10",
         8, False, WH)
    ], anchor=MSO_ANCHOR.MIDDLE)
    tb(slide, Inches(12.2), Inches(7.28), Inches(0.9), Inches(0.22), [
        (str(n), 9, True, WH, PP_ALIGN.RIGHT)
    ], anchor=MSO_ANCHOR.MIDDLE)


def title_bar(slide, title):
    rect(slide, Inches(0), Inches(0), W, Inches(0.72), G)
    tb(slide, Inches(0.32), Inches(0.12), Inches(12.7), Inches(0.52), [
        (title, 22, True, WH, PP_ALIGN.LEFT)
    ], anchor=MSO_ANCHOR.MIDDLE)


def style_table(table, header=True, col_fills=None, font=9):
    for r_i, row in enumerate(table.rows):
        for c_i, cell in enumerate(row.cells):
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for child in list(tcPr):
                if "mar" in child.tag:
                    tcPr.remove(child)
            for m, v in [("marL", "40000"), ("marR", "40000"), ("marT", "30000"), ("marB", "30000")]:
                el = etree.SubElement(tcPr, qn(f"a:{m}"))
                el.set("sz", v)
            fill = G if (header and r_i == 0) else (col_fills[c_i] if col_fills and r_i else (GP if r_i % 2 and not header else WH))
            if header and r_i == 0:
                fill = G
            elif col_fills and r_i > 0:
                fill = col_fills[c_i] if c_i < len(col_fills) else WH
            elif r_i % 2:
                fill = GP
            else:
                fill = WH
            solid = etree.SubElement(tcPr, qn("a:solidFill"))
            srgb = etree.SubElement(solid, qn("a:srgbClr"))
            srgb.set("val", rgb_hex(fill))
            for p in cell.text_frame.paragraphs:
                p.space_after = Pt(0)
                p.space_before = Pt(0)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(font)
                    run.font.bold = (r_i == 0 and header) or (c_i == 0 and r_i > 0)
                    run.font.color.rgb = WH if (header and r_i == 0) else GD if c_i == 0 else BK


def add_table(slide, l, t, w, h, data, font=9, header=True, col_w=None):
    rows, cols = len(data), len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, l, t, w, h)
    table = table_shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            table.columns[i].width = cw
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            set_run(r, str(val), font, bold=(i == 0 or j == 0), color=BK)
    style_table(table, header=header, font=font)
    return table


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # -------- 1 Title --------
    s = prs.slides.add_slide(blank)
    rect(s, Inches(0), Inches(0), W, Inches(1.85), G)
    tb(s, Inches(0.4), Inches(0.28), Inches(12.5), Inches(0.7), [
        ("Lumentum Holdings (NASDAQ: LITE)  |  Long Recommendation", 28, True, WH)
    ])
    tb(s, Inches(0.4), Inches(1.0), Inches(12.5), Inches(0.55), [
        ("UChicago AI-Enabled Investing Competition  ·  Round 1  ·  18 September 2026", 16, False, GL)
    ])
    tb(s, Inches(0.4), Inches(2.5), Inches(12.5), Inches(0.9), [
        ("Safety commitments coexist with similar capex. That is the entire thesis.", 22, True, G, PP_ALIGN.CENTER)
    ])
    tb(s, Inches(0.8), Inches(3.35), Inches(11.7), Inches(1.1), [
        ("Dario Amodei’s 12 September essay does not call for a hyperscaler capex cut. Prior safety regimes did not produce one. Among LITE, COHR and FN, Lumentum is the cleanest 3–12 month expression of continued optical spend.",
         15, False, GD, PP_ALIGN.CENTER)
    ])
    # three metric boxes
    boxes = [
        ("Pitched price (17 Sep close)", "$893.61"),
        ("12-month base target", "$1,003  |  +12%"),
        ("Bull / Bear (12-month P/E)", "$1,605  /  $331"),
    ]
    for i, (lab, val) in enumerate(boxes):
        left = Inches(0.7 + i * 4.15)
        rect(s, left, Inches(4.7), Inches(3.9), Inches(1.55), GP, G)
        tb(s, left + Inches(0.12), Inches(4.82), Inches(3.66), Inches(0.4), [(lab, 12, True, G, PP_ALIGN.CENTER)])
        tb(s, left + Inches(0.12), Inches(5.22), Inches(3.66), Inches(0.8), [(val, 22, True, GD, PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 1)

    # -------- 2 Exec summary --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "Executive Summary — LONG LITE, 3–12 months, because pacing ≠ capex pullback")
    # left dense pitch
    tb(s, Inches(0.28), Inches(0.82), Inches(8.55), Inches(6.35), [
        ("The PM’s question. Last week Dario Amodei published “We Must Pace the Frontier,” and investors read it as a signal that hyperscalers may pull back compute capex, hitting chip and infrastructure suppliers. We were asked to pick LONG or SHORT one of Lumentum, Coherent, or Fabrinet. Our answer is LONG LITE.", 12, False, BK),
        ("The variant view, in one sentence. Safety commitments coexist with similar capital expenditure. Pacing, as written, is about embedded evaluators, alignment work, and coordination — not a halt to training, cluster buildout, or optical interconnect. Amodei says explicitly that pacing “does not mean halting model training or technical progress.”", 12, False, BK),
        ("What we did this morning. We read the essay. We checked LITE’s tape around the weekend. We then walked a long list of prior safety events — the March 2023 FLI pause letter, Anthropic’s RSP, the Bletchley Declaration, the Seoul frontier commitments, the 2023 White House EO, and SB 1047 — and asked a single question: did infrastructure spend contract? It did not. Microsoft cash PP&E rose 58% then 45% across the Bletchley/Seoul windows. After the FLI letter, labs accelerated.", 12, False, BK),
        ("The tape this week is a narrative, not a revision. LITE fell 9.9% on Monday 14 September (close $835) and recovered to $919 by Wednesday. Frozen reference is Thursday’s $893.61 close. That is a one-session supply-chain scare. No hyperscaler cut a capex line. Buyer stocks (Meta, Alphabet, Microsoft) did not confirm a spend stop.", 12, False, BK),
        ("Why Lumentum, not the other two. LITE is the optical-content compounder inside the cluster — CPO lasers, ELS, NPO, OCS, 1.6T modules — already printing the ramp (Q4 revenue $1.01bn, +109% YoY; Q1 guide $1.25bn midpoint, 39.5–40.5% adj. OM, ahead of the target model). It is FCF-positive ($300m CFO−capex) with $1.1bn net cash and an NVIDIA multi-year purchase commitment plus $2bn investment. COHR funds growth with −$1.02bn cash conversion. FN is EMS with ~$4m of FCF after capex. If spend continues, LITE has the operating leverage. If the scare is wrong, LITE is the name that was mis-tagged as a capex victim.", 12, False, BK),
        ("How the trade gets paid. Next 3–12 months: Q1 print vs $1.225–1.275bn / $4.05–$4.35; subsequent quarters on product ramps and gross margin; hyperscaler capex commentary. Base 12-month value is 30× FY28 EPS $33.42 = $1,003. DCF ($545) is a conservative intrinsic cross-check, not the trading target — 72% terminal, SBC treated as cash, no multiple for the growth year. Bull is $1,605 if content/mix hits; bear is $331 if training orders are haircut. We are long because the historical base rate on “safety talk cuts capex” is approximately zero, and LITE is the best way to own that fact.", 12, True, GD),
    ])
    # right KPI stack
    kpis = [
        ("Recommendation", "LONG LITE"),
        ("Horizon", "3–12 months"),
        ("Entry (17 Sep)", "$893.61"),
        ("Base 12m PT", "$1,003 (+12%)"),
        ("FY27E / FY28E EPS", "$21.83 / $33.42"),
        ("Q1 FY27 guide", "$1.25bn  |  40% OM"),
        ("FY26 revenue", "$3.01bn  (+83%)"),
        ("Net cash", "$1.10bn"),
        ("CFO − capex", "+$300m"),
        ("NVIDIA", "Purchase + $2bn"),
    ]
    y = 0.82
    for lab, val in kpis:
        rect(s, Inches(8.95), Inches(y), Inches(4.1), Inches(0.58), WH, G)
        tb(s, Inches(9.05), Inches(y), Inches(1.55), Inches(0.58), [(lab, 10, True, G)], anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(10.55), Inches(y), Inches(2.4), Inches(0.58), [(val, 12, True, GD)], anchor=MSO_ANCHOR.MIDDLE)
        y += 0.62
    footer(s, 2)

    # -------- 3 Thesis summary (Sprouts page 4 clone, one thesis) --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "Investment Thesis Summary")
    tb(s, Inches(0.32), Inches(0.82), Inches(12.7), Inches(0.42), [
        ("The market is treating a safety essay as a capex cut. History, the essay’s own text, and this week’s tape all reject that mapping. LITE is the vehicle.", 14, True, G)
    ])
    # thesis table
    add_table(s, Inches(0.32), Inches(1.28), Inches(12.7), Inches(1.55), [
        ["", "The claim"],
        ["Single thesis", "Safety commitments coexist with similar (or higher) infrastructure capex. Pacing, RSPs, summits and pause letters have not historically reduced cluster spend. Therefore the optical interconnect cycle is intact over a 3–12 month horizon."],
        ["Why it pays in 3–12m", "Q1–Q4 FY27 prints and hyperscaler capex commentary either confirm the spend or they do not. We do not need a 10-year DCF to be right; we need the next four quarters of optical content not to be cancelled."],
    ], font=12, col_w=[Inches(2.2), Inches(10.5)])
    add_table(s, Inches(0.32), Inches(2.95), Inches(12.7), Inches(1.35), [
        ["Information edge", "1) Primary read of Amodei (12 Sep 2026) vs the capex-cut inference  2) LITE daily tape 8–17 Sep from stockanalysis  3) Event set of prior safety regimes vs Microsoft cash PP&E and lab behavior  4) Three-name screen on FCF, net cash, NVIDIA contract, and in-rack optics mix  5) Dated DCF / 12-month P/E model in Lumentum_Investment_Model.xlsx"],
    ], font=12, col_w=[Inches(2.2), Inches(10.5)])
    # three cases
    cases = [
        ("Bear case", "$331", "−63%", "Training-order haircut 10–20% from Q2, GM fade, 25× FY28 EPS $13.22. This is the PM’s world if pacing actually bites demand."),
        ("Base case", "$1,003", "+12%", "No training haircut. Q1 inside guide. FY28 EPS $33.42 at 30x. Capex/revenue 15% to 6%. This is spend continues, multiple normalizes."),
        ("Bull case", "$1,605", "+80%", "Volume/mix beats, GM to 59%, 35× FY28 EPS $45.86. CPO/OCS/1.6T layer in as management described in August."),
    ]
    for i, (name, px, ret, note) in enumerate(cases):
        left = Inches(0.32 + i * 4.3)
        rect(s, left, Inches(4.45), Inches(4.1), Inches(2.55), GP, G)
        rect(s, left, Inches(4.45), Inches(4.1), Inches(0.42), G)
        tb(s, left, Inches(4.45), Inches(4.1), Inches(0.42), [(name, 14, True, WH, PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
        tb(s, left + Inches(0.1), Inches(4.95), Inches(3.9), Inches(0.7), [(px, 26, True, GD, PP_ALIGN.CENTER)])
        tb(s, left + Inches(0.1), Inches(5.55), Inches(3.9), Inches(0.32), [(ret + " vs $893.61", 14, True, G, PP_ALIGN.CENTER)])
        tb(s, left + Inches(0.12), Inches(5.92), Inches(3.86), Inches(0.95), [(note, 11, False, BK, PP_ALIGN.LEFT)])
    footer(s, 3)

    # -------- 4 Evidence --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "Thesis evidence — Dario this week, and a long list of prior safety events")
    tb(s, Inches(0.32), Inches(0.80), Inches(7.3), Inches(0.38), [
        ("This week: a one-day optical scare that reversed. Not a capex print.", 13, True, G)
    ])
    s.shapes.add_picture(str(CHARTS / "lite_tape.png"), Inches(0.28), Inches(1.18), Inches(7.35), Inches(3.15))
    tb(s, Inches(0.32), Inches(4.38), Inches(7.3), Inches(2.55), [
        ("12 Sep (Sat): essay published. 14 Sep (Mon): LITE −9.9% to $835, volume 4.9m. 16 Sep: $919. 17 Sep freeze: $893.61.", 12, False, BK),
        ("Amodei’s own words reject the mapping. Pacing is extra time for alignment and embedded third-party evaluators. Step 1 is unilateral at Anthropic. Steps 2–3 need industry and government coordination that has not occurred this week.", 12, False, BK),
        ("Buyer vs supplier split. Secondary tape on 14 Sep had Meta/Alphabet up and Microsoft roughly flat while chips sold off. That is the market tagging the supply chain with a demand hole the buyers did not confirm. We will not pretend LITE “did not pull back” for one session. We will say the pullback did not persist and did not come with a capex cut.", 12, False, BK),
    ])
    add_table(s, Inches(7.75), Inches(0.82), Inches(5.28), Inches(6.15), [
        ["Prior safety event", "What followed on spend"],
        ["Mar 2023  FLI 6-month pause letter", "No pause. Labs accelerated. WIRED six-month check: development sped up, including signatories."],
        ["Sep 2023  Anthropic RSP", "Conditional training pause only if safeguards fail. Nov 2024 Amazon/Trainium partnership added $4bn (total $8bn). Financing, not a cut."],
        ["Oct 2023  White House EO 14110", "Reporting and evaluations. Not a capex freeze. Later revoked Jan 2025."],
        ["Nov 2023  Bletchley Declaration", "MSFT cash PP&E $28.1bn FY23 → $44.5bn FY24 (+58%). Window overlaps the summit."],
        ["May 2024  Seoul frontier commitments", "MSFT cash PP&E $44.5bn → $64.6bn FY25 (+45%). Extreme-risk language, rising spend."],
        ["Sep 2024  CA SB 1047", "Passed legislature, vetoed. No observed hyperscaler capex stop."],
        ["Feb 2025  Paris AI Action Summit", "Continuation of the Bletchley/Seoul process, not a build halt."],
        ["Sep 2026  Amodei essay", "No hyperscaler capex revision as of 18 Sep. LITE scare reversed in two sessions."],
    ], font=8, col_w=[Inches(2.05), Inches(3.23)])
    footer(s, 4)

    # -------- 5 Why LITE --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "Why Lumentum is the vehicle for this specific thesis")
    tb(s, Inches(0.32), Inches(0.80), Inches(12.7), Inches(0.38), [
        ("If spend continues, own the content lever with visible earnings and a clean balance sheet: LITE, not COHR or FN.", 13, True, G)
    ])
    add_table(s, Inches(0.28), Inches(1.18), Inches(12.75), Inches(2.55), [
        ["Screen (FY26 / as of 17 Sep)", "LITE", "COHR", "FN"],
        ["Price", "$893.61", "$295.98", "$380.92"],
        ["FY26 revenue", "$3,014m  (+83%)", "$7,118m", "$4,641m"],
        ["FY26 CFO − cash capex", "+$300m", "−$1,023m", "+$4m"],
        ["Net cash / (debt)", "+$1,101m", "−$1,235m", "+$875m"],
        ["FY27 adj. EPS consensus", "$21.67", "$9.42", "$18.16"],
        ["P / FY27 EPS", "41.2x", "31.4x", "21.0x"],
        ["What you actually own", "Lasers, modules, OCS inside the rack", "Broad materials / lasers; capex sponge", "EMS / assembly; thin FCF"],
        ["NVIDIA (Mar 2026)", "Purchase commitment + $2bn", "Partnership (no public $2bn)", "Not the content owner"],
    ], font=10, col_w=[Inches(2.7), Inches(3.35), Inches(3.35), Inches(3.35)])
    s.shapes.add_picture(str(CHARTS / "peer_fcf.png"), Inches(0.28), Inches(3.90), Inches(6.3), Inches(3.05))
    tb(s, Inches(6.7), Inches(3.90), Inches(6.3), Inches(3.05), [
        ("Why this thesis is not “own any AI supplier.”", 13, True, G),
        ("FN converts almost none of the cycle to free cash after capex and is one step removed from in-rack optics (CPO/ELS/NPO). A continued-spend world still leaves FN as a capacity contractor with Pillar Two tax and rising capex in the 10-K.", 12, False, BK),
        ("COHR is larger but funded the cycle with >$1bn of negative cash conversion. If spend continues, they still need the capital markets. If spend pauses, they are the most exposed.", 12, False, BK),
        ("LITE already ran through the target model “more than a quarter ahead of schedule” (Hurlston, 11 Aug). Q4 components $649m / systems $357m. Ultra-high-power CPO lasers, first ELS order, NPO engagements, 1.6T, OCS — that is incremental content per cluster, which is exactly what survives if training pacing is real but inference and networking are not cut.", 12, False, BK),
        ("We are paying 41x FY27 EPS for that. The 12-month trade is 30x FY28 $33.42, i.e. earnings growth, not multiple expansion.", 12, True, GD),
    ])
    footer(s, 5)

    # -------- 6 Financials --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "Financials — the ramp is already in the reported numbers")
    add_table(s, Inches(0.28), Inches(0.85), Inches(12.75), Inches(2.85), [
        ["Reported (USD m, except EPS)", "FY25", "FY26", "Q4 FY26", "Q1 FY27 guide"],
        ["Revenue", "1,645", "3,014", "1,006", "1,225–1,275"],
        ["Components / Systems", "1,116 / 529", "2,006 / 1,008", "649 / 357", "Not guided by SKU"],
        ["Adj. gross margin", "34.7%", "46.0%", "50.4%", "—"],
        ["Adj. operating margin", "9.7%", "29.8%", "36.6%", "39.5–40.5%"],
        ["Adj. EPS", "$2.06", "$8.67", "$3.23", "$4.05–$4.35"],
        ["CFO / cash capex", "126 / 231", "751 / 451", "—", "—"],
        ["Cash & ST investments", "877", "2,738", "2,738", "—"],
        ["Debt carrying value", "2,573", "1,637", "1,637", "—"],
    ], font=11, col_w=[Inches(2.7), Inches(2.0), Inches(2.35), Inches(2.7), Inches(3.0)])
    add_table(s, Inches(0.28), Inches(3.82), Inches(12.75), Inches(3.15), [
        ["Base model (ours, not consensus)", "Q1:27", "Q2:27", "Q3:27", "Q4:27", "FY27", "FY28"],
        ["Revenue ($m)", "1,250", "1,420", "1,620", "1,810", "6,101", "8,961"],
        ["Seq. growth", "+24% vs Q4", "+14%", "+14%", "+12%", "+102% YoY", "+47%"],
        ["Adj. OM", "40.0%", "42.2%", "43.8%", "44.7%", "42.9%", "45.8%"],
        ["Adj. EPS", "$4.17", "$5.00", "$5.92", "$6.75", "$21.83", "$33.42"],
        ["Economic uFCF ($m)", "—", "—", "—", "—", "940", "1,716"],
        ["Check vs company", "Inside guide", "", "", "", "EPS vs cons. +0.8%", "30x = $1,003"]
    ], font=11, col_w=[Inches(2.35), Inches(1.73), Inches(1.73), Inches(1.73), Inches(1.73), Inches(1.73), Inches(1.75)])
    footer(s, 6)

    # -------- 7 Valuation --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "Valuation & model assumptions — 12-month P/E is the target; DCF is the cross-check")
    tb(s, Inches(0.32), Inches(0.80), Inches(12.7), Inches(0.48), [
        ("We do not force the DCF and the trading target to agree. Base 12-month value is 30× FY28 EPS $33.42 = $1,003 (+12%). DCF at fiscal year-end is $545 (−39%). The long is an earnings-growth trade over 3–12 months, not a 10.5% WACC deep-value claim.", 13, True, G)
    ])
    add_table(s, Inches(0.28), Inches(1.32), Inches(6.4), Inches(3.55), [
        ["Assumption (Base, editable in xlsx)", "Value"],
        ["WACC / terminal g / terminal ROIC", "10.5% / 3.5% / 20%"],
        ["Diluted-share proxy", "102m (near Q4 adj. 101.1m)"],
        ["Q1 volume (comp / systems)", "+23.2% / +26.1% vs Q4"],
        ["FY28 volume (comp / systems)", "+45% / +50%"],
        ["Price/mix (Base)", "0%  (Bull +0.5%/qtr)"],
        ["Training-order haircut (Base / Bear)", "0% / 10–20% from Q2"],
        ["Adj. GM path FY27 qtrs", "52.5 → 55.5%"],
        ["Cash capex / sales FY27–33", "15% → 6%"],
        ["Inventory days FY27–33", "115 → 90"],
        ["SBC treated as economic cash", "3% of sales, deducted in DCF"],
        ["12-month exit P/E B/Base/Bull", "25× / 30× / 35×"],
        ["DSO / DPO / other NWC", "50 / 70 / 3% of sales"],
    ], font=10, col_w=[Inches(3.55), Inches(2.85)])
    add_table(s, Inches(6.85), Inches(1.32), Inches(6.18), Inches(2.55), [
        ["DCF build (Base, $m except /sh)", ""],
        ["FY27–33 economic uFCF", "940 → 1,716 → … → 5,836"],
        ["PV of explicit FCF", "15,380"],
        ["PV of terminal (g/ROIC reinvestment)", "39,146  (72% of EV)"],
        ["Enterprise value", "54,526"],
        ["+ cash − debt", "+2,738 − 1,637"],
        ["Equity / DCF / share", "55,627  /  $545"],
        ["12m P/E value", "$1,003  (+12%)"],
    ], font=11, col_w=[Inches(3.55), Inches(2.63)])
    add_table(s, Inches(6.85), Inches(3.95), Inches(6.18), Inches(1.85), [
        ["WACC \\ g", "2.5%", "3.5%", "4.5%"],
        ["8.5%", "709", "786", "899"],
        ["10.5% (used)", "514", "545", "586"],
        ["12.5%", "398", "413", "430"],
    ], font=11, col_w=[Inches(1.7), Inches(1.49), Inches(1.49), Inches(1.50)])
    tb(s, Inches(0.32), Inches(5.00), Inches(6.4), Inches(1.95), [
        ("What is fact vs assumption. Q1 Base is calibrated to issued guidance. Later quarters, annual volume, GM fade, capex ratios, WACC, g, ROIC, exit multiples, 40%/35% training-exposure shares, and 102m shares are ours. Product units and ASPs are not disclosed. DCF uses FY-end cash/debt and full-year cash flows — no fake September stub. GAAP FY26 loss includes a $7.76bn noncash convert charge; excluded from operations.", 11, False, BK),
        ("Invalidate. Shipments miss while inventory days rise, or adj. margins weaken despite growth. That is the bear, and it is in the file as case 1.", 11, True, GD),
    ])
    footer(s, 7)

    # -------- 8 Risks --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "The other side of the trade — what would make us wrong")
    add_table(s, Inches(0.28), Inches(0.88), Inches(12.75), Inches(5.95), [
        ["Risk", "Why it matters in 3–12 months", "How we live with it"],
        ["Pacing becomes a real training-compute cap, not a process overlay", "Bear case haircuts 10–20% of exposed orders from Q2. FY28 EPS falls to ~$13; 25× → $331 (−63%).", "Essay text and 2023–25 history argue against this. First tell is hyperscaler capex lines, not op-eds. Position size assumes this tail exists."],
        ["LITE is not cheap on DCF", "$545 vs $894. Terminal is 72% of EV. A long multiple on FY28 can compress if rates or AI multiples break.", "We are not selling a DCF long. We are selling 3–12 month earnings delivery vs a false capex-cut narrative. If you need DCF upside, this is not your idea."],
        ["Execution / yield / capacity", "Management cites CPO, ELS, NPO, 1.6T, OCS as layering in. None of that is a disclosed SKU P&L. Greensboro laser output is mid-2028 — outside the holding period.", "Q1 is a hard check: $1.225–1.275bn and 39.5–40.5% OM. Inventory days are modeled down 115→90; if they go the other way, sell."],
        ["Customer concentration and cancellations", "NVIDIA commitment is non-exclusive, multi-year, not a public quarterly revenue floor. Cloud customers can slow modules without a press release.", "Thesis requires continued cluster spend, not a single customer. Two missed quarters vs guide ends the idea."],
        ["COHR/FN relative", "FN is cheaper on FY27 EPS (21x). If the judges want “cheap AI supplier,” we lose the relative argument.", "Cheap and FCF-less is not the thesis. We want the content owner that already prints 50% GM."],
        ["Share count / SBC / convert", "Adj. diluted shares rose 71m → 101m. DCF charges 3% SBC as cash and holds 102m flat. Converts already produced a $7.76bn GAAP loss.", "If dilution runs faster than SBC, EPS is overstated. Watch the 10-Q share count, not the press-release EPS."],
        ["We may be early, or the scare already faded", "By 17 Sep LITE had retraced most of Monday. 12% base upside is modest if the narrative is already dead.", "Then the remaining pay is FY28 compounding, not mean-reversion. Bull still exists if CPO/OCS is real. Modest base + convex bull is acceptable given ~0% historical hit rate on the bear."],
    ], font=10, col_w=[Inches(2.35), Inches(5.2), Inches(5.2)])
    footer(s, 8)

    # -------- 9 Catalysts --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "How the trade gets paid — catalysts over the next 3–12 months")
    tb(s, Inches(0.32), Inches(0.82), Inches(12.7), Inches(0.4), [
        ("The highest-weight catalyst is the next earnings print versus issued Q1 ranges. Everything else is confirmation that spend was not cut.", 13, True, G)
    ])
    add_table(s, Inches(0.28), Inches(1.28), Inches(12.75), Inches(4.35), [
        ["Window", "Event", "What must be true for us", "What kills the idea"],
        ["0–3 months", "Q1 FY27 results vs $1.225–1.275bn, 39.5–40.5% adj. OM, $4.05–$4.35 EPS", "Print inside/above; commentary that CPO/OCS/1.6T are still ramping; no capex warning from MSFT/GOOGL/AMZN/META/ORCL", "Miss on revenue or OM; “push-outs”; a hyperscaler capex cut"],
        ["3–6 months", "Q2 print + first full post-essay capex season", "Seq. growth still positive without a training-order hole; GM holds ≥54% in our Base", "Haircut language; inventory days up with shipments down"],
        ["6–12 months", "Q3/Q4 + FY28 visibility", "Street FY28 EPS moves toward ~$33; 30× is then $1,000 without multiple expansion", "FY28 EPS stuck near FY27; multiple compresses with no earnings catch-up"],
        ["Any time", "Another safety essay / regulation", "Same test as Sep 12: text vs capex line. We have the checklist.", "A coordinated, verifiable training-compute cap that shows up in orders"],
        ["Not in the 3–12m pay", "Greensboro laser factory (mid-2028 target)", "Ignore as earnings. Capex is in the ratios; do not add a plant NPV.", "If someone capitalizes it as 2026–27 EPS, they are wrong"],
    ], font=11, col_w=[Inches(1.55), Inches(3.15), Inches(4.15), Inches(3.9)])
    tb(s, Inches(0.32), Inches(5.75), Inches(12.7), Inches(1.25), [
        ("Positioning statement. This is a 3–12 month long, not a multi-year compounder pitch and not a DCF-gap close. We get paid if the market stops mapping safety headlines onto optical demand, and if Lumentum simply delivers the guide it already issued on 11 August. The information edge is the event study plus the three-name screen — not a private dataset.", 13, False, GD),
        ("Entry $893.61  ·  Base $1,003  ·  Bull $1,605  ·  Bear $331  ·  Horizon 3–12 months  ·  Invalidation: two misses vs guide or a real capex cut", 14, True, G),
    ])
    footer(s, 9)

    # -------- 10 AI log --------
    s = prs.slides.add_slide(blank)
    title_bar(s, "AI-use log, verification, and hallucinations — required disclosure")
    tb(s, Inches(0.32), Inches(0.80), Inches(12.7), Inches(0.4), [
        ("Tools: Cursor agent + web retrieval + local Excel model. AI retrieved sources, filled formulas, and drafted slides. Forecasts are authored assumptions. No management interviews. No claim of human verification that has not occurred.", 12, True, G)
    ])
    add_table(s, Inches(0.28), Inches(1.22), Inches(12.75), Inches(5.75), [
        ["Item", "Status"],
        ["LITE 8–17 Sep prices", "Pulled from stockanalysis.com/stocks/lite/history/ on 18 Sep 2026. 17 Sep close $893.61 used as freeze. Intraday 18 Sep not used."],
        ["Amodei essay", "Read locally from archived HTML. Quote on “does not mean halting model training” is from the original, not a summary model."],
        ["LITE FY26 / Q1 guide", "Company 11 Aug 2026 release, archived. GAAP $7.76bn convert loss excluded from operating forecasts, as in the 10-K/release."],
        ["COHR / FN comps", "From the model Actuals/Comps tabs (company releases + stockanalysis). Not a full three-statement model of those names."],
        ["Microsoft PP&E +58% / +45%", "Cited in the workbook Evidence tab to Microsoft IR cash-flow pages. Team should click through; we did not re-parse the 10-K HTML in this deck build."],
        ["FLI pause / Bletchley / Seoul / EO / SB 1047 / Paris", "Primary pages and contemporaneous reporting. These are overlapping political windows, not independent causal estimates. We do not claim a zero-impact regression coefficient."],
        ["NVIDIA $2bn + purchase commitment", "NVIDIA 2 Mar 2026 announcement. Non-exclusive; no quarterly revenue floor inferred."],
        ["Base DCF $545 and 12m $1,003", "From Lumentum_Investment_Model.xlsx (case = Base). Bear/Bull 12m P/E cases replicated from the Assumptions tab (case 1/3); only Base values were cached in the xlsx snapshot."],
        ["Training share 40% / 35%", "ASSUMPTION. Not disclosed. Bear haircut applies only to this exposed slice."],
        ["102m diluted shares", "Proxy near Q4 adj. 101.1m. Not a convert-by-convert model."],
        ["Secondary 14 Sep Meta/Alphabet/Microsoft tape", "From market wrap articles. Treat as color; LITE path above is the verified series."],
        ["Hallucinations caught / avoided", "Did not use live quotes mixed with the 17 Sep freeze. Did not treat consensus forecast pages as audited. Did not add Greensboro as FY27 revenue. Did not claim LITE had zero Monday drawdown. Did not invent team-member names. Did not pretend DCF and 12m target agree."],
        ["Still unresolved", "Training vs inference split; SKU units/ASPs; contract cancellation terms; exact current diluted shares; human review of this deck."],
    ], font=10, col_w=[Inches(3.3), Inches(9.45)])
    footer(s, 10)

    prs.save(OUT)
    print("wrote", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    build()
