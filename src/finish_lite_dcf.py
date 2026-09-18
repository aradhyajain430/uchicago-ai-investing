"""Build final editable valuation workbook and ten-slide DCF section."""
from lite_dcf_model import *
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name as C
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

validate()
NAVY='#142638';TEAL='#087F8C';PALE='#EAF3F4';AMBER='#FFF1D6';GRAY='#536575'
wb=xlsxwriter.Workbook(OUT/'LITE_Bull_Case_DCF.xlsx')
wb.set_calc_mode('auto')
wb.set_properties({'title':'LITE | 3–6 month fear-unwind valuation','author':'Competition team, AI-assisted'})
fmt={
 'title':wb.add_format({'bold':True,'font_size':20,'font_color':'white','bg_color':NAVY}),
 'head':wb.add_format({'bold':True,'font_color':'white','bg_color':TEAL,'text_wrap':True}),
 'text':wb.add_format({'font_size':10,'text_wrap':True,'valign':'top'}),
 'note':wb.add_format({'font_size':10,'font_color':GRAY,'text_wrap':True,'valign':'top'}),
 'num':wb.add_format({'num_format':'#,##0.0;(#,##0.0);–','font_size':10}),
 'pct':wb.add_format({'num_format':'0.0%;(0.0%);–','font_size':10}),
 'input':wb.add_format({'num_format':'#,##0.00','font_color':'#1558C0','bg_color':'#EFF4FF'}),
 'ipct':wb.add_format({'num_format':'0.0%','font_color':'#1558C0','bg_color':'#EFF4FF'}),
 'total':wb.add_format({'num_format':'#,##0.0;(#,##0.0)','bold':True,'top':1,'bg_color':PALE}),
 'usd':wb.add_format({'num_format':'$#,##0.00;($#,##0.00)','bold':True,'font_size':12}),
 'warn':wb.add_format({'text_wrap':True,'bg_color':AMBER,'valign':'top','font_color':'#734600'}),
}
names=['Summary','Assumptions','Segments','Bull','Base','Bear','WACC','Capital','Sensitivity','Reverse DCF','Fear Unwind','Comps','Bear Liquidity','Sources','AI and Q&A','Checks']
ws={n:wb.add_worksheet(n) for n in names}
for name,s in ws.items():
    s.hide_gridlines(2);s.freeze_panes(6,1);s.set_zoom(90)
    s.set_column('A:A',40);s.set_column('B:K',14)
    s.merge_range('A1:K2','LITE | '+name.upper(),fmt['title'])
    s.merge_range('A3:K3','USD millions / shares millions except prices. Blue = inputs; black = formulas. Thesis horizon: 3–6 months.',fmt['note'])
    s.set_landscape();s.set_paper(9);s.fit_to_pages(1,0)
    s.set_footer('&LAI-assisted; sources and assumptions disclosed&R&P / &N')

def put(s,cell,val,kind='num'):s.write(cell,val,fmt[kind])
def f(s,cell,expr,val,kind='num'):s.write_formula(cell,expr,fmt[kind],val)
def band(s,r,text):s.merge_range(r-1,0,r-1,10,text,fmt['head']);s.set_row(r-1,24)
def note(s,r,text,warn=False,height=45):s.merge_range(r-1,0,r-1,10,text,fmt['warn' if warn else 'note']);s.set_row(r-1,height)
def years(s,r=6):
    for i,y in enumerate(YEARS,1):put(s,f'{C(i)}{r}',f'FY{y}','head')

# Audited capital and explicit treasury-stock method.
s=ws['Capital'];band(s,5,'REPORTED JUNE 27 CAPITAL | ONE CONSISTENT SNAPSHOT [S2]')
cr={7:('Cash',CAP['cash']),8:('Short-term investments',CAP['investments']),9:('Nonconvertible term debt',CAP['other_debt']),
10:('Common shares',CAP['common']),11:('Preferred, one-for-one common equivalent',CAP['preferred']),12:('RSUs outstanding',CAP['rsu']),
13:('PSUs outstanding, target payout',CAP['psu']),14:('Options outstanding',CAP['options']),15:('Option exercise price',CAP['strike']),16:('Reference close: Sept 17 [S3]',PRICE),
19:('Minority interest claim',0),20:('Separate preferred senior claim',0),21:('Operating NWC, reported inputs',CAP['nwc'])}
for r,(label,v) in cr.items():put(s,f'A{r}',label,'text');put(s,f'B{r}',v,'input');s.set_row(r-1,28)
put(s,'A17','Option dilution: treasury method','text');f(s,'B17','=B14*MAX(1-B15/B16,0)',TSM)
put(s,'A18','Core diluted shares before notes','text');f(s,'B18','=SUM(B10:B13)+B17',CORE,'total')
band(s,23,'CONVERTIBLES | CASH PRINCIPAL + PREMIUM SHARES AT FAIR VALUE')
for c,t in [('A24','Maturity'),('B24','Principal'),('C24','Conversion price'),('D24','Underlying shares')]:put(s,c,t,'head')
for r,n in enumerate(NOTES,25):
    put(s,f'A{r}',n['year']);put(s,f'B{r}',n['p'],'input');put(s,f'C{r}',n['k'],'input');f(s,f'D{r}',f'=B{r}/C{r}',n['p']/n['k'])
f(s,'B29','=SUM(B25:B28)',sum(n['p'] for n in NOTES),'total');f(s,'D29','=SUM(D25:D28)',ALL_SHARES-CORE,'total')
put(s,'A31','Full economic as-converted shares','text');f(s,'B31','=B18+D29',ALL_SHARES,'total')
put(s,'A32','Cash less nonconvertible debt','text');f(s,'B32','=B7+B8-B9',NET_CASH,'total')
put(s,'A33','Total debt principal','text');f(s,'B33','=B9+B29',DEBT)
note(s,35,'Preferred has no senior liquidation/redemption claim; include its common-equivalent shares instead of deducting $2bn. Options use the treasury method at the dated reference price. RSUs/PSUs use full outstanding shares at target vesting; no unrecognized-compensation offset, a conservative simplification. Capped-call value excluded.',height=58)
note(s,37,'Cash-settlement bridge subtracts all principal and divides by core shares plus only the conversion premium shares. Out-of-money notes remain debt. June cash/shares/debt are not mixed with August common shares. Post-June transactions and cash generation are not fully rolled forward.',True,54)
note(s,39,'NWC = 520.3 AR + 691.6 inventory + 211.6 other current assets − 567.4 AP − 146.3 payroll − 64.9 accrued − 91.5 other current liabilities = 553.4. Leases remain operating costs; no second lease-debt deduction.',height=43)

# WACC is built from components, with a peer beta calculation.
s=ws['WACC'];band(s,5,'NORMALIZED OPERATING RISK | NOT CALIBRATED TO A DESIRED SHARE PRICE')
for j,t in enumerate(['Peer','Beta','Equity value','Debt','D/E','Tax proxy','Unlevered beta']):put(s,f'{C(j)}7',t,'head')
for r,p in enumerate(PEERS,8):
    for c,v in [('A',p['ticker']),('B',p['beta']),('C',p['market_equity']),('D',p['debt']),('F',.25)]:put(s,f'{c}{r}',v,'text' if c=='A' else 'ipct' if c=='F' else 'input')
    f(s,f'E{r}',f'=D{r}/C{r}',p['de'],'pct');f(s,f'G{r}',f'=B{r}/(1+(1-F{r})*E{r})',p['unlevered_beta'])
wrows={13:('Risk-free rate [S6]',RF),14:('ERP observed [S7]',ERP_OBS),15:('ERP selected',ERP),19:('Operating cash tax',TAX),22:('Debt credit spread assumption',.025),28:('Own observed LITE beta [S15]',OWN_BETA),30:('Terminal g',G),31:('Exit EBITDA multiple',EXIT)}
for r,(label,v) in wrows.items():put(s,f'A{r}',label,'text');put(s,f'B{r}',v,'input' if r in [28,31] else 'ipct')
wform={16:('Median peer unlevered beta','=MEDIAN(G8:G10)',PEER_BETA),17:('Equity-equivalent market value','=Capital!B31*Capital!B16',MARKET_EQUITY),18:('Debt weight numerator','=Capital!B9',WACC_DEBT),
20:('LITE relevered beta','=B16*(1+(1-B19)*B18/B17)',BETA),21:('Cost of equity','=B13+B20*B15',KE),23:('Pretax marginal cost of debt','=B13+B22',KD),
24:('Debt weight','=B18/(B17+B18)',DW),25:('Equity weight','=1-B24',1-DW),26:('COMPONENT WACC','=B21*B25+B23*(1-B19)*B24',WACC),29:('Own-beta WACC cross-check','=(B13+B28*B15)*B25+B23*(1-B19)*B24',OWN_WACC)}
for r,(label,expr,v) in wform.items():put(s,f'A{r}',label,'text');f(s,f'B{r}',expr,v,'pct' if r in [21,23,24,25,26,29] else 'num')
note(s,34,'Beta source: public vendor 5-year betas, COHR / MRVL / FN [S8–S10]. Unlever with market D/E and a common 25% peer tax proxy; relever at LITE’s economic capital structure. All scenarios use the same WACC. No country premium or liquidity discount is added.',height=48)
note(s,36,'In-money notes are equity equivalents in the primary valuation; therefore WACC financing debt is the $92.8m term debt. Their contractual cash-principal liquidity requirement remains explicit. Borrowing cost is Treasury + assumed 250bp credit spread, not the misleading 0.375–1.5% convertible coupons.',height=50)
note(s,38,'Conservatism: selected ERP 4.50% exceeds the observed 4.14%; peer beta exceeds LITE’s own 1.54. A 3% terminal growth rate is below the roughly 4% long-run nominal benchmark (Fed June 2026: 2% real growth + 2% inflation). An independently sourced WACC is not proof the stock trades at fair value.',True,55)

s=ws['Segments'];years(s);band(s,5,'FORMER END-MARKET BUILD | FY26 ALLOCATION IS AN ESTIMATE')
labels={7:'Base Cloud / Networking',8:'Base Industrial Tech',9:'Base total revenue',10:'EML revenue uplift vs Base Cloud',11:'1.6T revenue uplift vs Base Cloud',12:'Incremental OCS revenue vs Base',13:'Internal sourcing margin increment',14:'Bull Cloud / Networking',15:'Bull Industrial Tech',16:'Bull total revenue',17:'Industrial growth',18:'Base total revenue growth'}
for r,t in labels.items():put(s,f'A{r}',t,'text')
for i,y in enumerate(YEARS,1):
    cc=C(i);prev=C(i-1);idx=i-1
    for r,v in [(10,EML[idx]),(11,MIX[idx]),(12,OCS[idx]),(13,SOURCING[idx])]:put(s,f'{cc}{r}',v,'input' if r==12 else 'ipct')
    gi=IND[idx]/(IND26 if i==1 else IND[idx-1])-1;put(s,f'{cc}17',gi,'ipct')
    f(s,f'{cc}8',f'=$B$23*(1+{cc}17)' if i==1 else f'={prev}8*(1+{cc}17)',IND[idx])
    if i<3:put(s,f'{cc}9',BASE_TOTAL[idx],'input');f(s,f'{cc}18',f'={cc}9/'+('$B$24' if i==1 else f'{prev}9')+'-1',BASE_TOTAL[idx]/(3014 if i==1 else BASE_TOTAL[idx-1])-1,'pct')
    else:put(s,f'{cc}18',BASE_TOTAL[idx]/BASE_TOTAL[idx-1]-1,'ipct');f(s,f'{cc}9',f'={prev}9*(1+{cc}18)',BASE_TOTAL[idx])
    for r,expr,v in [(7,f'={cc}9-{cc}8',BASE_CLOUD[idx]),(14,f'={cc}7*(1+{cc}10+{cc}11)+{cc}12',CASES['Bull']['cloud'][idx]),(15,f'={cc}8',IND[idx]),(16,f'={cc}14+{cc}15',RESULTS['Bull']['rows'][idx]['revenue'])]:f(s,f'{cc}{r}',expr,v,'total' if r==16 else 'num')
for r,label,v in [(21,'FY25 reported Industrial [S5]',234.2),(22,'Estimated FY26 Industrial growth',.05),(24,'FY26 reported total [S1]',3014)]:put(s,f'A{r}',label,'text');put(s,f'B{r}',v,'ipct' if r==22 else 'input')
put(s,'A23','FY26 Industrial estimate','text');f(s,'B23','=B21*(1+B22)',IND26)
put(s,'A25','FY26 Cloud estimate','text');f(s,'B25','=B24-B23',3014-IND26)
note(s,27,'FY25 reported Cloud / Networking and Industrial Tech were $1,410.8m / $234.2m. FY26 switched to one reporting segment with Components / Systems product disclosure ($2,005.6m / $1,008.4m). The former-segment estimates above reconcile to the reported consolidated total; they are not company-disclosed FY26 end markets.',True,58)
note(s,29,'Bull excess revenue is attributed only to LITE-specific EML share/content, 1.6T mix and incremental OCS. The percentages are revenue-impact proxies, not asserted actual market shares or disclosed mix percentages. Internal sourcing changes margin only. No internal component sale is counted twice.',height=52)

# Scenario statements, terminal values and capital bridge.
for name,case in CASES.items():
    s=ws[name];r=RESULTS[name];years(s);band(s,5,'FY27–31 RAMP + FY32–36 NORMALIZATION | '+name.upper())
    labels={8:'Cloud revenue growth',9:'Industrial revenue growth',10:'Adjusted operating margin',11:'SBC and related payroll / sales',12:'Cash tax rate',13:'Depreciation / sales',14:'Capex / sales',15:'Operating NWC / sales',18:'Cloud / Networking revenue',19:'Industrial Tech revenue',20:'Total revenue',21:'Revenue growth',23:'Adjusted EBIT',24:'Less recurring SBC / payroll',25:'Economic EBIT',26:'Cash operating taxes',27:'NOPAT',28:'Add D&A: depreciation only',29:'Less capex',30:'Operating NWC balance',31:'Less change in NWC',32:'UNLEVERED FREE CASH FLOW',33:'Economic EBIT margin',34:'Economic EBITDA',36:'Remaining-year fraction',37:'Discount years',38:'Discount factor',39:'PV of uFCF'}
    for rr,label in labels.items():put(s,f'A{rr}',label,'text')
    for i,d in enumerate(r['rows'],1):
        cc=C(i);prev=C(i-1);idx=i-1
        for rr,key in [(8,'cloud'),(9,'industrial')]:
            prior=(3014-IND26 if key=='cloud' else IND26) if i==1 else r['rows'][idx-1][key]
            put(s,f'{cc}{rr}',d[key]/prior-1,'ipct')
        if name=='Bull':f(s,f'{cc}10',f'=Base!{cc}10+Segments!{cc}13',d['margin'],'pct')
        else:put(s,f'{cc}10',d['margin'],'ipct')
        for rr,v in [(11,case['sbc'][idx]),(13,case['dep'][idx]),(14,case['capex'][idx]),(15,case['nwc'])]:put(s,f'{cc}{rr}',v,'ipct')
        f(s,f'{cc}12','=WACC!$B$19',TAX,'pct')
        for rr,key,sr in [(18,'cloud',14 if name=='Bull' else 7),(19,'industrial',15 if name=='Bull' else 8)]:
            if name!='Bear':f(s,f'{cc}{rr}',f'=Segments!{cc}{sr}',d[key])
            elif i==1:put(s,f'{cc}{rr}',d[key],'input')
            else:f(s,f'{cc}{rr}',f'={prev}{rr}*(1+{cc}{8 if rr==18 else 9})',d[key])
        equations={20:(f'={cc}18+{cc}19',d['revenue']),21:(f'={cc}20/'+('Segments!$B$24' if i==1 else f'{prev}20')+'-1',d['revenue']/(3014 if i==1 else r['rows'][idx-1]['revenue'])-1),
        23:(f'={cc}20*{cc}10',d['adj_ebit']),24:(f'={cc}20*{cc}11',d['sbc']),25:(f'={cc}23-{cc}24',d['ebit']),26:(f'={cc}25*{cc}12',d['tax']),27:(f'={cc}25-{cc}26',d['nopat']),28:(f'={cc}20*{cc}13',d['da']),29:(f'={cc}20*{cc}14',d['capex']),30:(f'={cc}20*{cc}15',d['nwc']),31:(f'={cc}30-'+('Capital!$B$21' if i==1 else f'{prev}30'),d['dnwc']),32:(f'={cc}27+{cc}28-{cc}29-{cc}31',d['fcf']),33:(f'={cc}25/{cc}20',d['ebit']/d['revenue']),34:(f'={cc}25+{cc}28',d['ebitda']),36:('=Summary!$B$25' if i==1 else '=1',d['fraction']),37:(f'=Summary!$B$25+{idx}',d['time']),38:(f'=1/(1+WACC!$B$26)^{cc}37',1/(1+WACC)**d['time']),39:(f'={cc}32*{cc}36*{cc}38',d['pv'])}
        for rr,(expr,v) in equations.items():f(s,f'{cc}{rr}',expr,v,'pct' if rr in [21,33,36] else 'total' if rr in [20,25,32,39] else 'num')
    band(s,41,'TERMINAL VALUE: GORDON GROWTH AND MATURE EXIT EBITDA MULTIPLE')
    terminal={43:('Terminal revenue','=K20*(1+WACC!B30)',r['terminal_revenue']),44:('Terminal NOPAT','=B43*(K10-K11)*(1-K12)',r['terminal_revenue']*(case['margin'][-1]-case['sbc'][-1])*(1-TAX)),45:('Terminal depreciation','=B43*K13',r['terminal_revenue']*case['dep'][-1]),46:('Terminal capex','=B43*K14',r['terminal_revenue']*case['capex'][-1]),47:('Terminal delta NWC','=B43*K15-K30',r['terminal_revenue']*case['nwc']-r['rows'][-1]['nwc']),48:('Normalized terminal uFCF','=B44+B45-B46-B47',r['terminal_fcf']),49:('Undiscounted terminal EV','=B48/(WACC!B26-WACC!B30)',r['tv']),50:('PV terminal value','=B49*K38',r['tvpv']),51:('PV explicit forecast','=SUM(B39:K39)',r['pv_explicit']),52:('ENTERPRISE VALUE','=B50+B51',r['ev']),53:('TV / EV','=B50/B52',r['tv_share']),54:('Gordon implied terminal EBITDA multiple','=B49/K34',r['implied_terminal_multiple'])}
    for rr,(label,expr,v) in terminal.items():put(s,f'A{rr}',label,'text');f(s,f'B{rr}',expr,v,'pct' if rr==53 else 'total' if rr in [48,52] else 'num')
    for cell,expr,val in [('D49','=K34*WACC!B31',r['exit_tv']),('D50','=D49*K38',r['exit_pv']),('D51','=B51',r['pv_explicit']),('D52','=D50+D51',r['exit_ev']),('D53','=D50/D52',r['exit_tv_share'])]:f(s,cell,expr,val,'pct' if cell=='D53' else 'num')
    put(s,'B42','Gordon','head');put(s,'D42','Exit method','head')
    band(s,56,'EV → EQUITY → PER SHARE | CONTRACTUAL CASH-PRINCIPAL PRESENTATION')
    bridge_labels={58:'Enterprise value',59:'Add cash and investments',60:'Less total debt principal',61:'Less minority / preferred senior claims',62:'Equity after all principal',63:'Core diluted shares',64:'Add net conversion premium shares',65:'Total net diluted shares',66:'VALUE PER SHARE',67:'Upside / downside vs reference'}
    for rr,label in bridge_labels.items():put(s,f'A{rr}',label,'text')
    for cc,method,ev in [('B','gordon',r['ev']),('D','exit',r['exit_ev'])]:
        b=r[method]
        # Five closed-form regimes avoid circularity from cash principal / premium shares.
        regcol='G' if cc=='B' else 'I';kcol='H' if cc=='B' else 'J'
        sorted_notes=sorted(enumerate(NOTES,25),key=lambda pair:pair[1]['k'])
        for n in range(5):
            rr=71+n;inactive=sorted_notes[n:];active=sorted_notes[:n]
            debt='+'.join(f'Capital!B{k}' for k,_ in inactive) or '0'
            sh='+'.join(f'Capital!D{k}' for k,_ in active) or '0'
            candidate=(ev+NET_CASH-sum(v['p'] for _,v in inactive))/(CORE+sum(v['p']/v['k'] for _,v in active))
            f(s,f'{regcol}{rr}',f'=({cc}52+Capital!B32-Capital!B19-Capital!B20-({debt}))/(Capital!B18+{sh})',candidate)
            if n<4:f(s,f'{kcol}{rr}',f'=Capital!C{sorted_notes[n][0]}',sorted_notes[n][1]['k'])
        value_formula=f'=IF({regcol}71<{kcol}71,{regcol}71,IF({regcol}72<{kcol}72,{regcol}72,IF({regcol}73<{kcol}73,{regcol}73,IF({regcol}74<{kcol}74,{regcol}74,{regcol}75))))'
        prem='+'.join(f'IF({cc}66>Capital!C{k},Capital!D{k}-Capital!B{k}/{cc}66,0)' for k in range(25,29))
        entries={58:(f'={cc}52',ev),59:('=Capital!B7+Capital!B8',CAP['cash']+CAP['investments']),60:('=Capital!B33',DEBT),61:('=Capital!B19+Capital!B20',0),62:(f'={cc}58+{cc}59-{cc}60-{cc}61',b['equity_after_all_principal']),63:('=Capital!B18',CORE),64:('='+prem,b['net_shares']-CORE),65:(f'={cc}63+{cc}64',b['net_shares']),66:(value_formula,b['value']),67:(f'={cc}66/Capital!B16-1',b['value']/PRICE-1)}
        for rr,(expr,v) in entries.items():f(s,f'{cc}{rr}',expr,v,'usd' if rr==66 else 'pct' if rr==67 else 'num')
    note(s,69,'D&A add-back excludes acquired amortization because adjusted EBIT already removes it. Recurring SBC remains an economic cost. Terminal adjusted margin is 25%, economic EBIT margin 20.5%; peak margins are not perpetual. FY27 cash flow is prorated uniformly to the remaining fiscal year.',height=52)
    s.set_row(69,8);s.print_area('A1:K69')

# Assumptions ledger.
s=ws['Assumptions'];s.set_column('A:A',31);s.set_column('B:B',30);s.set_column('C:C',76);s.set_column('D:D',33)
for j,t in enumerate(['Driver','Value / path','Logic / limitation','Source']):put(s,f'{C(j)}5',t,'head')
for r,row in enumerate(ASSUMPTIONS,6):
    for j,v in enumerate(row):put(s,f'{C(j)}{r}',v,'text')
    s.set_row(r-1,60)

# Sensitivity grids, formula-driven with terminal FCF normalized at each g.
s=ws['Sensitivity'];band(s,5,'BULL | WACC × TERMINAL GROWTH | $/SHARE')
RATES=[WACC-.02,WACC-.01,WACC,WACC+.01,WACC+.02];GS=[.02,.025,.03,.035,.04];MULTS=[12,14,16,18,20]
def share_formula(ev):
    sorted_notes=sorted(enumerate(NOTES,25),key=lambda pair:pair[1]['k'])
    candidates=[]
    for n in range(5):
        inactive=sorted_notes[n:];active=sorted_notes[:n]
        debt='+'.join(f'Capital!$B${k}' for k,_ in inactive) or '0';sh='+'.join(f'Capital!$D${k}' for k,_ in active) or '0'
        candidates.append(f'(({ev})+Capital!$B$32-Capital!$B$19-Capital!$B$20-({debt}))/(Capital!$B$18+{sh})')
    formula=candidates[-1]
    for n in range(3,-1,-1):formula=f'IF(({candidates[n]})<Capital!$C${sorted_notes[n][0]},({candidates[n]}),({formula}))'
    return '='+formula
for start,kind,cols in [(7,'g',GS),(18,'exit',MULTS)]:
    if kind=='exit':band(s,16,'BULL | WACC × EXIT EBITDA MULTIPLE | $/SHARE')
    for j,v in enumerate(cols,1):put(s,f'{C(j)}{start}',v,'ipct' if kind=='g' else 'input')
    for rr,w in enumerate(RATES,start+1):
        put(s,f'A{rr}',w,'ipct')
        for j,v in enumerate(cols,1):
            cc=C(j);pv='+'.join(f'Bull!{C(k)}32*Bull!{C(k)}36/(1+$A{rr})^Bull!{C(k)}37' for k in range(1,11))
            if kind=='g':
                tr=f'Bull!$K$20*(1+{cc}${start})'
                tf=f'(({tr})*((Bull!$K$10-Bull!$K$11)*(1-Bull!$K$12)+Bull!$K$13-Bull!$K$14)-(({tr})*Bull!$K$15-Bull!$K$30))'
                terminal=f'{tf}/($A{rr}-{cc}${start})/(1+$A{rr})^Bull!$K$37';vout=evaluate(CASES['Bull'],wacc=w,g=v)['gordon']['value']
            else:terminal=f'Bull!$K$34*{cc}${start}/(1+$A{rr})^Bull!$K$37';vout=evaluate(CASES['Bull'],wacc=w,exit_multiple=v)['exit']['value']
            f(s,f'{cc}{rr}',share_formula(pv+'+'+terminal),vout,'usd')
    s.conditional_format(f'B{start+1}:F{start+5}',{'type':'3_color_scale'})
note(s,26,'Central cells match Bull. Gordon normalizes terminal working capital at each g. Exit method applies the multiple to FY2036 economic EBITDA after recurring SBC. These are alternative methods, not values to add together. A 16× exit assumes materially more value than the Gordon-implied mature multiple.',True,54)

# Reverse DCF carries full formula forecasts; Goal Seek can update the solved parameters.
s=ws['Reverse DCF'];put(s,'A5','Solved FY28–31 Cloud CAGR','text');put(s,'B5',REV_GROWTH,'ipct');put(s,'A6','Solved FY28–31 adjusted margin','text');put(s,'B6',REV_MARGIN,'ipct')
for start,mode,res in [(9,'growth',REVERSE_G),(35,'margin',REVERSE_M)]:
    band(s,start,'MARKET-IMPLIED '+mode.upper()+' | OTHER BULL DRIVERS HELD FIXED');years(s,start+1)
    labels={2:'Cloud growth',3:'Cloud revenue',4:'Industrial revenue',5:'Total revenue',6:'Adjusted margin',7:'Economic EBIT',8:'NOPAT',9:'Depreciation',10:'Capex',11:'Operating NWC',12:'Change in NWC',13:'uFCF',14:'PV of uFCF'}
    for off,label in labels.items():put(s,f'A{start+off}',label,'text')
    for i,d in enumerate(res['rows'],1):
        cc=C(i);prev=C(i-1);idx=i-1
        growth=(d['cloud']/res['rows'][idx-1]['cloud']-1) if i>1 else d['cloud']/(3014-IND26)-1
        growth_expr=f'=$B$5' if i<=5 else f'=$B$5+(WACC!$B$30-$B$5)*{i-5}/5'
        if i==1 or mode=='margin':growth_expr=f'=Bull!{cc}8'
        f(s,f'{cc}{start+2}',growth_expr,growth,'pct')
        f(s,f'{cc}{start+3}',f'=Bull!{cc}18' if i==1 or mode=='margin' else f'={prev}{start+3}*(1+{cc}{start+2})',d['cloud'])
        f(s,f'{cc}{start+4}',f'=Bull!{cc}19',d['industrial'])
        f(s,f'{cc}{start+5}',f'={cc}{start+3}+{cc}{start+4}',d['revenue'])
        marg_expr=f'=Bull!{cc}10' if i==1 or mode=='growth' else '=$B$6' if i<=5 else f'=$B$6+(Bull!$K$10-$B$6)*{i-5}/5'
        f(s,f'{cc}{start+6}',marg_expr,d['margin'],'pct')
        eqs={7:(f'={cc}{start+5}*({cc}{start+6}-Bull!{cc}11)',d['ebit']),8:(f'={cc}{start+7}*(1-Bull!{cc}12)',d['nopat']),9:(f'={cc}{start+5}*Bull!{cc}13',d['da']),10:(f'={cc}{start+5}*Bull!{cc}14',d['capex']),11:(f'={cc}{start+5}*Bull!{cc}15',d['nwc']),12:(f'={cc}{start+11}-'+('Capital!$B$21' if i==1 else f'{prev}{start+11}'),d['dnwc']),13:(f'={cc}{start+8}+{cc}{start+9}-{cc}{start+10}-{cc}{start+12}',d['fcf']),14:(f'={cc}{start+13}*Bull!{cc}36*Bull!{cc}38',d['pv'])}
        for off,(expr,v) in eqs.items():f(s,f'{cc}{start+off}',expr,v)
    rr=start+17;tr=f'K{start+5}*(1+WACC!B30)'
    tf=f'({tr})*((K{start+6}-Bull!K11)*(1-Bull!K12)+Bull!K13-Bull!K14)-(({tr})*Bull!K15-K{start+11})'
    put(s,f'A{rr}','Implied enterprise value','text');f(s,f'B{rr}',f'=SUM(B{start+14}:K{start+14})+({tf})/(WACC!B26-WACC!B30)*Bull!K38',res['ev'])
    put(s,f'A{rr+1}','Implied share value','text');f(s,f'B{rr+1}',share_formula(f'B{rr}'),res['gordon']['value'],'usd')
    put(s,f'A{rr+2}','Difference vs market price','text');f(s,f'B{rr+2}',f'=B{rr+1}-Capital!B16',0)
note(s,60,'Reverse growth: FY27 fixed; solve constant Cloud CAGR for FY28–31, then fade it linearly to 3% by FY36. Reverse margin: hold Bull sales fixed; solve FY28–31 adjusted margin, then fade to 25%. If the solved margin exceeds 100%, the margin-only solution is economically infeasible.',True,56)
note(s,62,'Solved inputs are Python bisection outputs. After editing assumptions in Excel, use Goal Seek: set B28 to 0 by changing B5 (growth), or B54 to 0 by changing B6 (margin); alternatively rerun the build. These are conditional break-even paths, not a claim that all investors share one forecast.',height=49)
band(s,64,'CAPEX REALITY CHECK | REPORTED SPENDING SUPPORTS CONTINUITY, NOT ANY VALUATION')
note(s,65,'Meta: CY2026 capex guide $130–145bn incl. finance-lease principal [S13]. Amazon: TTM net PPE purchases $169.007bn, +64% [S14]. These are broad, overlapping infrastructure indicators with different periods; do not add them or equate them to an optical TAM. LITE FY26 capex $451.3m; the model funds its own ramp explicitly.',height=58)
put(s,'A68','Bull FY27 capex','text');f(s,'B68','=Bull!B29',RESULTS['Bull']['rows'][0]['capex'])
put(s,'A69','Reverse growth FY31 capex','text');f(s,'B69','=F19',REVERSE_G['rows'][4]['capex'])

s=ws['Fear Unwind'];band(s,5,'THESIS: FEAR DISCOUNT REVERSES WITHIN 3–6 MONTHS IF EXECUTION HOLDS')
for rr,label,v in [(7,'Pre-event close: Sept 11 [S3]',PRE_EVENT),(8,'Event close: Sept 14 [S3]',EVENT),(9,'Reference close: Sept 17 [S3]',PRICE),(10,'Analyst average, 12-month horizon [S4]',ANALYST_AVG)]:put(s,f'A{rr}',label,'text');put(s,f'B{rr}',v,'input')
for rr,label,expr,v in [(12,'Event decline','=B8/B7-1',EVENT/PRE_EVENT-1),(13,'Drop recovered by reference close','=(B9-B8)/(B7-B8)',(PRICE-EVENT)/(PRE_EVENT-EVENT)),(14,'Remaining pure unwind upside','=B7/B9-1',PRE_EVENT/PRICE-1),(15,'Analyst-average upside','=B10/B9-1',ANALYST_AVG/PRICE-1),(16,'Extra beyond prior price','=B10/B7-1',ANALYST_AVG/PRE_EVENT-1)]:put(s,f'A{rr}',label,'text');f(s,f'B{rr}',expr,v,'pct')
band(s,19,'CONDITIONAL PRICE-IMPLIED DISCOUNT RATE | NOT THE COMPONENT WACC')
for rr,label,v in [(21,'Pre-event price implied WACC',IMPLIED_PRE),(22,'Shock-close implied WACC',IMPLIED_SHOCK),(23,'Reference-close implied WACC',IMPLIED_CURRENT),(24,'Analyst-average implied WACC',IMPLIED_ANALYST)]:put(s,f'A{rr}',label,'text');put(s,f'B{rr}',v,'ipct')
note(s,27,'Holding all Bull cash flows fixed, these rates solve the Gordon DCF to each observed price/benchmark. Their difference isolates a possible risk-discount interpretation; it does not establish that the essay caused the move. The low absolute price-implied rates conflict with the independently built WACC.',True,55)
note(s,29,'Catalysts over 3–6 months: next 1–2 reports confirm revenue guidance, OCS / 1.6T shipments and capacity conversion, allowing a perceived risk discount to shrink. No causal event study or consensus revision time series was available. The $1,149 vendor target is a 12-month benchmark, not a demonstrated 3–6 month DCF fair value.',height=54)

s=ws['Comps'];band(s,5,'VALUATION CROSS-CHECKS | DATED CLOSES, MIXED BUSINESS MODELS')
for j,t in enumerate(['Peer','Close','Beta','EV / TTM EBITDA','Forward P/E','LITE implied value*']):put(s,f'{C(j)}7',t,'head')
for rr,p in enumerate(PEERS,8):
    vals=[p['ticker'],p['price'],p['beta'],p['ev_ebitda'],p['forward_pe'],p['forward_pe']*21.67]
    for j,v in enumerate(vals):put(s,f'{C(j)}{rr}',v,'text' if j==0 else 'input')
put(s,'A13','LITE FY27 consensus adjusted EPS','text');put(s,'B13',21.67,'input')
for rr,p in enumerate(PEERS,8):f(s,f'F{rr}',f'=E{rr}*$B$13',p['forward_pe']*21.67,'usd')
for rr,label,v in [(15,'Analyst low',ANALYST_LOW),(16,'Analyst average',ANALYST_AVG),(17,'Analyst high',ANALYST_HIGH)]:put(s,f'A{rr}',label,'text');put(s,f'B{rr}',v,'input')
note(s,20,'*Peer forward P/E × LITE FY27 consensus adjusted EPS is an indicative relative-value check. Vendor forward P/E may use NTM rather than LITE fiscal-year EPS. COHR, MRVL and FN differ in growth, margins and capital intensity; no automatic peer-median price target is adopted.',True,53)
note(s,22,'Peer EV / TTM EBITDA is reconstructed from Sept 17 closes, public shares, debt, cash and EBITDA. These are vendor common-equity EV proxies; complex peer preferred interests/leases may not be fully captured. The terminal 16× is a mature-state judgment, not the current 24–74× trading range.',height=50)

s=ws['Bear Liquidity'];years(s);band(s,5,'BEAR IS OPERATING-SURVIVABLE | NO NEW DEBT OR EQUITY REQUIRED IN THIS SIMPLE TEST')
for rr,label in [(8,'Opening cash / investments'),(9,'Economic uFCF, stub-adjusted'),(10,'Cash principal repayments'),(11,'After-tax interest stress'),(12,'Closing cash / investments'),(13,'Remaining debt principal')]:put(s,f'A{rr}',label,'text')
for i,r in enumerate(BEAR_LIQ,1):
    cc=C(i);prev=C(i-1);opening=CAP['cash']+CAP['investments'] if i==1 else BEAR_LIQ[i-2]['cash']
    f(s,f'{cc}8','=Capital!B7+Capital!B8' if i==1 else f'={prev}12',opening)
    f(s,f'{cc}9',f'=Bear!{cc}32*Bear!{cc}36',RESULTS['Bear']['rows'][i-1]['fcf']*RESULTS['Bear']['rows'][i-1]['fraction'])
    put(s,f'{cc}10',r['repayment'],'input')
    f(s,f'{cc}11',f'='+('Capital!$B$33' if i==1 else f'{prev}13')+f'*WACC!$B$23*(1-WACC!$B$19)*Bear!{cc}36',r['interest'])
    f(s,f'{cc}12',f'={cc}8+{cc}9-{cc}10-{cc}11',r['cash'],'total')
    f(s,f'{cc}13','='+('Capital!$B$33' if i==1 else f'{prev}13')+f'-{cc}10',r['debt'])
note(s,16,'All remaining note principal is cash-settled by modeled maturities; 2026 notes are in FY27, 2029 calendar maturity is conservatively paid in FY29. Term-loan repayments are approximated and total $92.8m. Premium conversion value settles in shares. Interest is stressed at marginal borrowing cost, above actual convertible coupons.',height=58)
note(s,18,'Economic FCF already charges recurring SBC, making this a conservative cash proxy. No buybacks, acquisitions or dividends. This annual model cannot establish intra-quarter liquidity or covenant compliance. All-note immediate principal stress leaves $1,091.3m before operating cash flows; restricted/required bank deposits may reduce available liquidity.',True,52)

s=ws['Summary'];band(s,5,'3–6 MONTH THESIS | SENTIMENT RECOVERY REQUIRES INTACT EARNINGS')
note(s,6,'The team expects a fear-driven discount to unwind over the next 1–2 earnings reports as AI infrastructure spending and optical shipments continue. The independent DCF below tests the thesis; it does not force a bullish target.',height=42)
put(s,'A8','Reference close: Sept 17','text');f(s,'B8','=Capital!B16',PRICE,'usd')
for j,name in enumerate(['Bear','Base','Bull'],1):
    cc=C(j);r=RESULTS[name];put(s,f'{cc}10',name,'head')
    for rr,label,ref,v in [(11,'Gordon value / share','B66',r['gordon']['value']),(12,'Gordon upside / downside','B67',r['gordon']['value']/PRICE-1),(13,'Exit value / share','D66',r['exit']['value']),(14,'Exit upside / downside','D67',r['exit']['value']/PRICE-1),(15,'FY27 revenue','B20',r['rows'][0]['revenue']),(16,'FY27 uFCF','B32',r['rows'][0]['fcf']),(17,'FY31 revenue','F20',r['rows'][4]['revenue']),(18,'Terminal adjusted margin','K10',r['rows'][-1]['margin']),(19,'Gordon TV / EV','B53',r['tv_share'])]:
        put(s,f'A{rr}',label,'text');f(s,f'{cc}{rr}',f'={name}!{ref}',v,'usd' if rr in [11,13] else 'pct' if rr in [12,14,18,19] else 'num')
note(s,21,'Interpretation: the operating Bull case does not support the reference share price under the conservative component WACC and mature-margin fade. A 3–6 month rebound can occur despite this; do not label a sentiment/analyst recovery target as DCF-derived intrinsic value.',True,51)
put(s,'A24','Valuation date','text');put(s,'B24','2026-09-18','text');put(s,'A25','FY27 remaining-year fraction','text');f(s,'B25','=(DATE(2027,6,26)-DATE(2026,9,18))/365',STUB,'pct')
put(s,'A27','Pure unwind to pre-event close','text');f(s,'B27',"='Fear Unwind'!B7",PRE_EVENT,'usd')
put(s,'A28','Pure unwind upside','text');f(s,'B28',"='Fear Unwind'!B14",PRE_EVENT/PRICE-1,'pct')
put(s,'A29','Analyst average, 12-month benchmark','text');f(s,'B29','=Comps!B16',ANALYST_AVG,'usd')
note(s,32,'Navigation: Assumptions and Segments show each driver; Bull / Base / Bear contain the complete uFCF build; WACC and Capital show all components; Sensitivity includes both grids; Reverse DCF and Fear Unwind separate market-implied fundamentals from sentiment; Bear Liquidity tests survival.',height=48)
note(s,34,'Formula outputs are cached and recalculate in Excel. Reverse solved inputs and price-implied discount rates require rerunning Python or Goal Seek after changes. June balance-sheet data and a prorated FY27 stub are disclosed simplifications. Slides are a snapshot of this build.',height=48)
s.set_column('B:D',20)

s=ws['Sources'];s.set_column('A:A',7);s.set_column('B:B',32);s.set_column('C:C',20);s.set_column('D:D',65);s.set_column('E:E',70)
for j,t in enumerate(['ID','Document','Date','URL','Use / limitation']):put(s,f'{C(j)}5',t,'head')
for rr,source in enumerate(SOURCES,6):
    for j,key in enumerate(['id','title','date','url','usage']):put(s,f'{C(j)}{rr}',source[key],'text')
    s.set_row(rr-1,70)
rr=6+len(SOURCES);put(s,f'A{rr}','S17','text');put(s,f'B{rr}','Fed June 2026 projections','text');put(s,f'C{rr}','2026-06-17','text');put(s,f'D{rr}','https://www.federalreserve.gov/monetarypolicy/fomcprojtabl20260617.htm','text');put(s,f'E{rr}','Long-run median real GDP 2%, inflation 2%; roughly 4% nominal benchmark.','text');s.set_row(rr-1,60)

QA=[
 ('Thesis ownership','User selected LITE, a Bull operating case and a 3–6 month fear-discount reversal thesis. A negative DCF result is retained rather than changing WACC to hit a target.'),
 ('Above-consensus','Bull FY27/FY28 sales exceed published $6.32bn/$9.61bn estimates only through the explicitly modeled EML / 1.6T / OCS revenue increments. Internal sourcing raises margin. Product-level consensus is unavailable, so these are above-base assumptions, not verified product-consensus beats.'),
 ('Beyond FY28','Public later-year consensus was unavailable. Ramp/fade, reinvestment and mature margin are independent assumptions. It would be false to confirm that every non-operating assumption matches consensus.'),
 ('WACC and g','Same conservative WACC/g across cases: peer beta exceeds own beta; ERP above observed; g 3% below ~4% nominal benchmark. These are independent estimates, not price-calibrated rates. The 16× exit multiple is a separate judgment and is more generous than Gordon.'),
 ('Fear is a hypothesis','The September 14 drop is observed; its attribution to the essay and unchanged post-event consensus are not established by a causal study or revision history. The current reference has already recovered part of that drop.'),
 ('Mature margin','Peak adjusted margin 42.5% falls to 25%; economic EBIT is another 4.5 percentage points lower after SBC. COHR 21.8% and MRVL 36.6% adjusted margins bracket the assumption but are not identical businesses or a mature-peer consensus.'),
 ('Accounting and sources','Company/SEC accounts anchor reported values. Vendor estimates and betas are labeled. Reconstructed former segments are estimates; no actual product share, OCS revenue, or internal sourcing percentage is invented.'),
 ('AI use','AI retrieved and checked sources, modeled cash flows, built Excel and slides, and ran arithmetic/formula checks. Human direction set thesis/company/horizon. Human verification is not represented as complete.'),
 ('Corrections','Prior FN recommendation and preliminary LITE model are superseded. Revised model includes component WACC, terminal margin fade, both terminal methods, options treasury method, product uplift drivers and reverse DCF. The earlier softer-WACC draft must not be used.'),
 ('Capital limits','Existing awards use full RSU/PSU counts; options treasury method uses entry reference price. June snapshot and later conversion requests are disclosed. Preferred is common-equivalent; capped-call benefits excluded. No double deduction of full convertible principal plus full conversion dilution.'),
]
s=ws['AI and Q&A']
for rr,(label,text) in enumerate(QA,5):put(s,f'A{rr}',label,'head');s.merge_range(f'B{rr}:K{rr}',text,fmt['text']);s.set_row(rr-1,66)

s=ws['Checks'];checks=[]
for name,r in RESULTS.items():
    checks.extend([(f'{name}: cash flow identity',f'={name}!B32-({name}!B27+{name}!B28-{name}!B29-{name}!B31)',0),
      (f'{name}: Gordon equity bridge',f'={name}!B62/{name}!B65-{name}!B66',0),
      (f'{name}: exit equity bridge',f'={name}!D62/{name}!D65-{name}!D66',0),
      (f'{name}: terminal margin fades',f'={name}!K10-0.25',0)])
checks.extend([('Sensitivity Gordon center','=Sensitivity!D10-Bull!B66',0),('Sensitivity exit center','=Sensitivity!D21-Bull!D66',0),('Reverse growth matches price',"='Reverse DCF'!B27-Capital!B16",0),('Reverse margin matches price',"='Reverse DCF'!B53-Capital!B16",0),('WACC greater than g','=IF(WACC!B26>WACC!B30,0,1)',0),('Bear annual liquidity positive',"=IF(MIN('Bear Liquidity'!B12:K12)>0,0,1)",0)])
for rr,(label,expr,v) in enumerate(checks,6):put(s,f'A{rr}',label,'text');f(s,f'B{rr}',expr,v);f(s,f'C{rr}',f'=IF(ABS(B{rr})<0.000001,"PASS","FAIL")','PASS','text');s.set_row(rr-1,26)
note(s,27,'Checks validate arithmetic and identities, not forecast realism. Reverse-margin arithmetic can match price while requiring an impossible >100% margin; that is an analytical finding, not a model pass on feasibility.',True,47)
wb.close()

# Final independent cached-data validation.
data=openpyxl.load_workbook(OUT/'LITE_Bull_Case_DCF.xlsx',data_only=True)
formulas=openpyxl.load_workbook(OUT/'LITE_Bull_Case_DCF.xlsx',data_only=False)
count=0
for sheet in formulas:
    for row in sheet:
        for cell in row:
            if cell.data_type=='f':
                count+=1;value=data[sheet.title][cell.coordinate].value
                assert value is not None,(sheet.title,cell.coordinate)
                assert not (isinstance(value,str) and value.startswith(('#REF!','#DIV/0!','#VALUE!','#NAME?','#NUM!'))),(sheet.title,cell.coordinate,value)
for name,r in RESULTS.items():assert abs(data[name]['B66'].value-r['gordon']['value'])<1e-7

# Presentation: editable native tables and shapes, ten slides limited to valuation.
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def rgb(h):return RGBColor.from_string(h.lstrip('#'))
def text(sl,x,y,w,h,t,size=18,color=NAVY,bold=False):
    box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=Inches(.015)
    for i,line in enumerate(str(t).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.font.bold=bold;p.space_after=Pt(5)
    return box
def rect(sl,x,y,w,h,color):
    sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return sh
def slide(title,subtitle,sources,notes):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);rect(sl,0,0,13.333,.16,TEAL)
    text(sl,.5,.34,12.3,.65,title,26,bold=True);text(sl,.5,1.04,12.3,.5,subtitle,14,GRAY)
    rect(sl,.5,6.79,12.3,.015,'#CDD9DF');text(sl,.5,6.91,11.7,.38,'Sources: '+sources+'  |  AI: source extraction, model and slides; team review pending.',8,GRAY)
    text(sl,12.38,6.88,.45,.35,len(prs.slides),10,GRAY)
    sl.notes_slide.notes_text_frame.text=notes+'\n\nSource URLs:\n'+'\n'.join(s['id']+' '+s['url'] for s in SOURCES if s['id'] in sources.split(', '))
    return sl
def table(sl,x,y,width,height,headers,rows,colweights=None,size=14):
    sh=sl.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(width),Inches(height));tb=sh.table
    weights=colweights or [1]*len(headers)
    for j,w in enumerate(weights):tb.columns[j].width=Inches(width*w/sum(weights))
    for i,row in enumerate([headers]+rows):
        for j,v in enumerate(row):
            cell=tb.cell(i,j);cell.text=str(v);cell.margin_left=Inches(.09);cell.margin_right=Inches(.06);cell.margin_top=Inches(.05)
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(TEAL if i==0 else PALE if i%2 else '#FFFFFF')
            for p in cell.text_frame.paragraphs:p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=i==0;p.font.color.rgb=rgb('#FFFFFF' if i==0 else NAVY)
    return sh
def dollar(v):return f'${v:,.0f}'
def million(v):return f'{v:,.0f}'
def pct(v):return f'{v:.1%}'
b=RESULTS['Bull'];base=RESULTS['Base'];bear=RESULTS['Bear']

sl=slide('A 3–6 month fear unwind needs earnings to hold','Team thesis: safety-related fears fade as the next 1–2 reports confirm optical demand and execution.','S1, S3, S4, S16','This is the user-directed thesis, not a causal conclusion. Distinguish the observed event-price recovery from additional upside to a 12-month analyst benchmark. Use the dated reference close consistently; do not claim it is live. The independent DCF is below the current price under the required conservative assumptions.')
table(sl,.55,1.8,12.2,2.6,['Valuation reference','Price','Vs. Sept 17 close','Interpretation'],[
 ['Sept 14 shock close',dollar(EVENT),pct(EVENT/PRICE-1),'Observed price; essay causality not established'],
 ['Sept 17 reference close',dollar(PRICE),'—','Frozen entry comparison'],
 ['Pre-event close: Sept 11',dollar(PRE_EVENT),pct(PRE_EVENT/PRICE-1),'Pure price-recovery benchmark'],
 ['Analyst average target',dollar(ANALYST_AVG),pct(ANALYST_AVG/PRICE-1),'12-month cross-check; needs more than pure unwind']], [2.3,1,1.3,4],15)
text(sl,.65,4.78,12,1.25,f'{pct((PRICE-EVENT)/(PRE_EVENT-EVENT))} of the event-day dollar drop had already recovered by the reference close.\nA $1,149 target requires an additional {pct(ANALYST_AVG/PRE_EVENT-1)} beyond the pre-event price.',20,bold=True)

sl=slide('Operating upside is explicit; the ramp and margins fade','Five years capture the ramp, then five years normalize growth and margins before terminal value.','S1, S2, S4, S5, S11, S12','All drivers are in the Assumptions sheet. FY26 former-segment allocation is estimated because LITE reorganized to one segment and reports Components / Systems products. EML and 1.6T percentages are revenue uplift proxies, not reported market share. OCS is incremental to Base, not total disclosed OCS revenue. Tax and valuation assumptions do not improve in Bull.')
table(sl,.55,1.67,12.2,4.7,['Driver','Bull assumption / path','Source or logic'],[
 ['Revenue by end market','Cloud + Industrial; FY27 '+million(b['rows'][0]['revenue'])+'m','Base FY27/28: $6.32bn / $9.61bn [S4]; former split estimated'],
 ['EML / 1.6T / OCS','EML +1–3%; mix +2–5%; OCS +$150m → $1.49bn','LITE-specific incremental revenue; explicit team estimates'],
 ['Adjusted EBIT margin','40.75% FY27 → 42.5% peak → 25% terminal','Sourcing adds up to 150bp, then fades; peer margin bracket'],
 ['SBC / cash tax','4.5% of sales / 20%','SBC remains an economic cost; tax above FY26 adjusted rate'],
 ['D&A / capex / NWC','D&A 4→5.5%; capex 13→6%; NWC 18%','Depreciation only; capex exceeds terminal depreciation'],
 ['Long-run growth / exit','3% g / 16× economic EBITDA','Below ~4% nominal GDP; mature multiple, not current peer peak'],
 ['Capital / timing','Treasury options; June capital; FY27 stub','Cash principal + net conversion shares; no double counting']], [1.6,3.1,4.5],13)

sl=slide('Full unlevered cash-flow build: FY2027–2031','USD millions. Economic EBIT retains recurring stock compensation; acquired amortization is not added twice.','S1, S2, S4','Forecast outputs are model assumptions. D&A means depreciation-only add-back because the adjusted EBIT starting point already removes acquired amortization. Every projected year, including FY32–36, is visible in Excel. FY27 is a full-year operating forecast; its DCF cash flow is then multiplied by the remaining-year fraction.')
metrics=[('Cloud / Networking','cloud',1),('Industrial Tech','industrial',1),('Revenue','revenue',1),('Adjusted EBIT','adj_ebit',1),('Less SBC / payroll','sbc',-1),('Economic EBIT','ebit',1),('Less cash tax','tax',-1),('NOPAT','nopat',1),('Add D&A','da',1),('Less capex','capex',-1),('Less ΔNWC','dnwc',-1),('Unlevered FCF','fcf',1)]
rows=[[label]+[million(d[key]*sign) for d in b['rows'][:5]] for label,key,sign in metrics]
table(sl,.55,1.65,12.2,4.87,['USD m']+[str(y) for y in YEARS[:5]],rows,[2.5,1,1,1,1,1],13)

sl=slide('WACC comes from risk components, not the price target',f'Primary WACC {WACC:.2%}; identical across Bear / Base / Bull. Own-beta cross-check: {OWN_WACC:.2%}.','S6, S7, S8, S9, S10, S15','Market-value peer debt/equity used for unlevering; 25% common peer tax proxy. Relever at LITE economic capital structure. In-money notes are treated as equity equivalents in WACC and valuation, while contractual cash principal is included in the liquidity test. Cost of debt is marginal risk-free plus spread, not the low convertible coupon.')
table(sl,.55,1.7,7.5,4.5,['Component','Value','Basis'],[
 ['Risk-free rate',f'{RF:.2%}','Sept 17 US 10-year Treasury'],['Equity risk premium',f'{ERP:.2%}',f'Above {ERP_OBS:.2%} observed implied ERP'],['Peer unlevered beta',f'{PEER_BETA:.2f}','Median COHR / MRVL / FN'],['Relevered LITE beta',f'{BETA:.2f}','Market weights; economic capital structure'],['Cost of equity',f'{KE:.2%}','Rf + beta × ERP'],['Pretax cost of debt',f'{KD:.2%}','Rf + assumed 250bp credit spread'],['Debt / equity weights',f'{DW:.2%} / {1-DW:.2%}','In-money notes treated as equity equivalents'],['WACC',f'{WACC:.2%}','Ke × E/V + Kd × (1−tax) × D/V']], [2,1.4,3.8],14)
table(sl,8.35,1.7,4.4,2.25,['Peer','Raw beta','Unlevered'],[[p['ticker'],f"{p['beta']:.2f}",f"{p['unlevered_beta']:.2f}"] for p in PEERS],size=14)
text(sl,8.4,4.3,4.25,1.8,'Conservative choices:\nERP above observed; peer beta above LITE’s 1.54; terminal growth 3%.\nNo lower discount rate in Bull.',17)

sl=slide('Terminal value uses a mature margin, under both methods','FY2036 adjusted EBIT margin 25%; economic EBIT margin 20.5% after recurring SBC.','S1, S2, S11, S12','Gordon normalizes FY2037 working-capital investment at 3% growth. Exit applies 16 times FY2036 economic EBITDA after SBC, not non-GAAP EBITDA before SBC. Both terminal values are discounted to September 18. The exit method is more generous than Gordon; do not average them automatically.')
table(sl,.55,1.7,8.0,4.45,['Bull terminal / DCF metric','Gordon growth','Exit multiple'],[
 ['Long-run assumption','3.0% growth','16.0× EBITDA'],['Normalized terminal FCF / EBITDA',million(b['terminal_fcf'])+'m',million(b['rows'][-1]['ebitda'])+'m'],['Terminal value, undiscounted',million(b['tv'])+'m',million(b['exit_tv'])+'m'],['PV terminal value',million(b['tvpv'])+'m',million(b['exit_pv'])+'m'],['PV explicit cash flows',million(b['pv_explicit'])+'m',million(b['pv_explicit'])+'m'],['Enterprise value',million(b['ev'])+'m',million(b['exit_ev'])+'m'],['TV AS % OF EV',pct(b['tv_share']),pct(b['exit_tv_share'])]], [2.7,1.6,1.6],15)
text(sl,8.9,1.95,3.8,.7,'No perpetual peak margin',21,bold=True)
text(sl,8.9,2.8,3.8,2.7,f'42.5% peak → 25% terminal.\n\nPeer context: COHR 21.8%; MRVL 36.6% adjusted margins.\n\nGordon implies only {b["implied_terminal_multiple"]:.1f}× terminal EBITDA.',17)

sl=slide('The per-share bridge exposes the valuation gap','Cash principal is deducted; only net conversion-premium shares enter the denominator.','S2, S3','Preferred participates one-for-one and has no senior liquidation claim; include shares, not a separate $2bn deduction. Options use 0.2m shares and $8.10 strike via treasury method at reference price. Full RSUs 2m and PSUs 0.8m included at target payout. Capped-call benefit excluded. Both cases below are standalone DCF values, not a forced recovery target.')
table(sl,.55,1.7,8.1,4.55,['USD m, except shares / prices','Gordon','Exit'],[
 ['Enterprise value',million(b['ev']),million(b['exit_ev'])],['Add cash / short-term investments',million(CAP['cash']+CAP['investments']),million(CAP['cash']+CAP['investments'])],['Less debt principal',million(-DEBT),million(-DEBT)],['Less minority / senior preferred','0 / 0','0 / 0'],['Equity after cash principal',million(b['gordon']['equity_after_all_principal']),million(b['exit']['equity_after_all_principal'])],['Core + net conversion shares',f'{b["gordon"]["net_shares"]:.2f}m',f'{b["exit"]["net_shares"]:.2f}m'],['Value per share',dollar(b['gordon']['value']),dollar(b['exit']['value'])],['Upside / downside vs '+dollar(PRICE),pct(b['gordon']['value']/PRICE-1),pct(b['exit']['value']/PRICE-1)]],[3,1.4,1.4],14)
rect(sl,9,1.85,3.75,4.28,AMBER);text(sl,9.18,2.05,3.38,3.85,'Interpretation\n\nThe operating Bull case is not enough to support today’s price under these conservative valuation assumptions.\n\nA sentiment rebound is possible; it is not a DCF-derived target.',18)

sl=slide('Sensitivity: neither terminal convention hides the risk','Bull cash flows; blue-axis inputs are editable in Excel. Values shown are dollars per share.','S1, S2, S6, S7','Central row uses exact component WACC. Gordon grid recalculates terminal FCF for each g; exit grid applies mature EBITDA multiples. Both use the full capital bridge including conversion regimes. These values are alternatives, not additive.')
headers=['WACC / g']+[f'{g:.1%}' for g in GS]
table(sl,.55,1.95,6.0,3.15,headers,[[f'{w:.2%}']+[dollar(evaluate(CASES['Bull'],wacc=w,g=g)['gordon']['value']) for g in GS] for w in RATES],[1.3,1,1,1,1,1],12)
headers=['WACC / exit']+[f'{v}×' for v in MULTS]
table(sl,6.8,1.95,6.0,3.15,headers,[[f'{w:.2%}']+[dollar(evaluate(CASES['Bull'],wacc=w,exit_multiple=v)['exit']['value']) for v in MULTS] for w in RATES],[1.3,1,1,1,1,1],12)
text(sl,.6,1.56,6,.3,'Gordon growth',16,bold=True);text(sl,6.85,1.56,6,.3,'Exit EBITDA multiple',16,bold=True)
text(sl,.7,5.52,11.8,.7,'Discount-rate normalization alone must be justified independently. The model does not reduce WACC merely to reproduce an analyst price target.',19,bold=True)

sl=slide('Cross-checks distinguish intrinsic value from a recovery trade','The analyst average is a 12-month benchmark; our proposed catalyst window is 3–6 months.','S3, S4, S8, S9, S10','Football field is a range display, not a probability distribution. Peer P/E range uses vendor forward multiples times LITE FY27 consensus adjusted EPS $21.67; NTM/fiscal-period mismatch limits comparability. Peer EV/EBITDA values are current trailing vendor proxies, not terminal assumptions. Analysts may use different growth, margins and discount rates.')
table(sl,.55,1.72,5.1,2.3,['Peer','EV / EBITDA','Fwd P/E'],[[p['ticker'],f'{p["ev_ebitda"]:.1f}×',f'{p["forward_pe"]:.1f}×'] for p in PEERS],size=15)
text(sl,.65,4.38,4.8,1.6,'Peer P/E × $21.67 LITE EPS:\n'+dollar(min(p['forward_pe']*21.67 for p in PEERS))+'–'+dollar(max(p['forward_pe']*21.67 for p in PEERS))+'\nIndicative, not a like-for-like target.',17)
fig,ax=plt.subplots(figsize=(7.1,4.25),dpi=160);fig.patch.set_facecolor('white')
ranges=[('Gordon: bear–bull',bear['gordon']['value'],b['gordon']['value']),('Exit: bear–bull',bear['exit']['value'],b['exit']['value']),('Peer P/E proxy',min(p['forward_pe']*21.67 for p in PEERS),max(p['forward_pe']*21.67 for p in PEERS)),('Analyst range',ANALYST_LOW,ANALYST_HIGH)]
for i,(label,lo,hi) in enumerate(ranges):ax.plot([lo,hi],[i,i],lw=13,color=TEAL,solid_capstyle='butt');ax.text(lo,i+.24,f'${lo:,.0f}',fontsize=9,ha='center');ax.text(hi,i+.24,f'${hi:,.0f}',fontsize=9,ha='center')
ax.scatter([ANALYST_AVG],[3],color='#EA9A36',s=65,zorder=5);ax.axvline(PRICE,color=NAVY,ls='--',lw=1.3,label=f'Ref. ${PRICE:.0f}')
ax.set_yticks(range(4),[x[0] for x in ranges]);ax.invert_yaxis();ax.set_xlim(0,1500);ax.set_ylim(3.65,-.5);ax.grid(axis='x',alpha=.15);ax.set_xlabel('LITE value per share ($)');ax.legend(loc='lower right',fontsize=9)
for side in ['top','right','left']:ax.spines[side].set_visible(False)
fig.tight_layout();chart=OUT/'LITE_Valuation_Football_Field.png';fig.savefig(chart,bbox_inches='tight');plt.close(fig)
sl.shapes.add_picture(str(chart),Inches(5.85),Inches(1.7),width=Inches(7.0),height=Inches(4.55))

sl=slide('Bear / Base / Bull: operating survival is not an equity floor','Only operating execution varies; common WACC, tax, terminal growth, terminal margin and exit multiple.','S1, S2, S4','Bear is not an insolvency scenario. Economic FCF remains positive and a simplified annual liquidity schedule pays all debt principal with no new financing. It does not establish intraperiod liquidity, covenant compliance or a share-price floor. A larger share-price downside is retained rather than artificially truncating the bear.')
rows=[]
for label,key,form in [('FY27 revenue','revenue',million),('FY27 uFCF','fcf',million),('FY31 revenue','rev31',million),('Gordon value','gv',dollar),('Exit value','xv',dollar),('Gordon return','return',pct)]:
    vals=[]
    for n in ['Bear','Base','Bull']:
        r=RESULTS[n];v=r['rows'][0][key] if key in ['revenue','fcf'] else r['rows'][4]['revenue'] if key=='rev31' else r['gordon']['value'] if key=='gv' else r['exit']['value'] if key=='xv' else r['gordon']['value']/PRICE-1;vals.append(form(v))
    rows.append([label]+vals)
table(sl,.55,1.75,8.0,3.65,['USD m except per share','Bear','Base','Bull'],rows,[2.5,1,1,1],15)
rect(sl,8.95,1.8,3.8,3.7,PALE);text(sl,9.15,2.0,3.4,3.3,f'Bear cash remains positive\n\nMinimum modeled year-end cash / investments:\n${min(r["cash"] for r in BEAR_LIQ)/1000:.2f}bn\n\nAll principal funded; no buybacks or acquisitions assumed.',18)
text(sl,.7,5.76,11.8,.6,'Invalidate the trade if orders slip, utilization weakens or the next guidance cannot support continuing execution.',19,bold=True)

sl=slide('Reverse DCF: what must be true at today’s price?','Conditional market-implied paths expose the gap between a near-term fear unwind and long-term valuation.','S1, S2, S3, S13, S14','Growth solve fixes FY27 Bull sales, solves FY28–31 Cloud CAGR and fades that growth to 3% over FY32–36. Margins use Bull path ending at 25%. Margin solve fixes Bull revenue and solves FY28–31 margin, fading to 25% thereafter. Capex flexes with sales. This is not an assertion of a single investor consensus. Price-implied rates are diagnostic and differ from independent WACC. Meta guidance and Amazon TTM spending periods differ; do not add them.')
table(sl,.55,1.7,7.65,3.8,['At '+dollar(PRICE),'Required outcome','Meaning'],[
 ['Growth-only solve',pct(REV_GROWTH)+' Cloud CAGR','FY28–31; fade to 3% by FY36'],
 ['Implied FY31 revenue','$'+f'{REVERSE_G["rows"][4]["revenue"]/1000:.1f}bn','vs Bull $'+f'{b["rows"][4]["revenue"]/1000:.1f}bn'],
 ['Implied FY31 capex','$'+f'{REVERSE_G["rows"][4]["capex"]/1000:.1f}bn','Explicit funding burden, not costless growth'],
 ['Margin-only solve',pct(REV_MARGIN),'Above 100%: economically infeasible'],
 ['Price-implied discount rate',pct(IMPLIED_CURRENT),'Below independent '+pct(WACC)+' WACC']], [2.1,1.8,2.8],14)
text(sl,8.58,1.85,4.05,.7,'Capex supports continuity',20,bold=True)
text(sl,8.58,2.65,4.05,2.1,'Meta: $130–145bn CY2026 guide.\nAmazon: $169bn TTM net PPE, +64%.\n\nThese support spending today; they do not prove the market-implied decade of LITE growth.',17)
rect(sl,.6,5.72,12.1,.67,AMBER);text(sl,.75,5.81,11.8,.5,'Defend the 3–6 month rebound as a conditional sentiment trade; the conservative DCF does not certify today’s price as cheap.',17,bold=True)
prs.save(OUT/'LITE_DCF_Slides.pptx')

summary={'valuation_date':'2026-09-18','horizon':'3–6 months','reference_price':PRICE,'wacc':WACC,'terminal_growth':G,'exit_multiple':EXIT,'cases':RESULTS,'reverse_growth':REV_GROWTH,'reverse_margin':REV_MARGIN,'bear_liquidity':BEAR_LIQ,'formula_cells':count,'validation':'Python arithmetic / cache checks PASS; Excel recalculation pending'}
(OUT/'LITE_Model_Outputs.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
qa='# LITE valuation: presenter Q&A\n\n'+ '\n\n'.join('**'+a+'**\n\n'+b for a,b in QA)
qa+='\n\n## Headline outputs\n\n'+f'Bull Gordon ${b["gordon"]["value"]:,.0f}; Bull exit ${b["exit"]["value"]:,.0f}; reference ${PRICE:,.2f}. Component WACC {WACC:.2%}. Pure pre-event recovery ${PRE_EVENT:,.2f} / {PRE_EVENT/PRICE-1:.1%}. Analyst mean ${ANALYST_AVG:,} is a 12-month external benchmark.\n'
(OUT/'LITE_DCF_QA.md').write_text(qa,encoding='utf-8')
print(json.dumps({'workbook':'LITE_Bull_Case_DCF.xlsx','slides':'LITE_DCF_Slides.pptx','slide_count':len(prs.slides),'formula_cells':count,'bull_gordon':b['gordon']['value'],'bull_exit':b['exit']['value'],'validation':'PASS'},indent=2))
