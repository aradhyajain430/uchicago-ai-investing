"""Build the final Lumentum pitch deck in a UChicago-maroon Sprouts-style template."""
from __future__ import annotations

from pathlib import Path
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "final_slideshow"
CHARTS = OUT / "charts"
OUT.mkdir(exist_ok=True)
CHARTS.mkdir(exist_ok=True)

# University of Chicago-inspired palette; intentionally no Sprouts green.
MAROON = "800000"
RED = "B13F3F"
LIGHT_RED = "E7C7C7"
PALE = "F5ECEC"
CREAM = "FBF9F6"
GRAY = "EFEFEF"
MID = "777777"
DARK = "241F20"
WHITE = "FFFFFF"
BLACK = "000000"
GOLD = "C8A951"


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def make_charts() -> None:
    plt.rcParams.update({"font.family": "DejaVu Serif", "axes.edgecolor": "#B7B7B7"})

    # Price reaction around the September essay.
    labels = ["Sep 8", "Sep 9", "Sep 10", "Sep 11", "Sep 14", "Sep 15", "Sep 16", "Sep 17"]
    prices = [978.54, 988.98, 935.70, 927.03, 835.03, 838.96, 919.40, 893.61]
    fig, ax = plt.subplots(figsize=(8.2, 3.3), dpi=180)
    ax.plot(labels, prices, color="#800000", lw=2.7, marker="o", ms=5)
    ax.axvspan(3.5, 4.5, color="#E7C7C7", alpha=.85)
    ax.axhline(prices[-1], color="#C8A951", ls="--", lw=1)
    ax.annotate("Sep 14: -9.9%", (4, 835.03), (1.7, 812), color="#800000", fontsize=9,
                arrowprops=dict(arrowstyle="->", color="#800000", lw=.9))
    ax.annotate("Sep 16: $919", (6, 919.40), (5.3, 975), color="#241F20", fontsize=9,
                arrowprops=dict(arrowstyle="->", color="#241F20", lw=.9))
    ax.set_ylim(790, 1020)
    ax.set_ylabel("LITE close ($)", fontsize=9)
    ax.grid(axis="y", color="#E7E1DE")
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(CHARTS / "price_reaction.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Safety commitments and Microsoft cash PP&E - overlapping, not causal.
    yrs = [2023, 2024, 2025, 2026]
    capex = [28.1, 44.5, 64.6, 115.9]
    fig, ax = plt.subplots(figsize=(7.4, 3.2), dpi=180)
    bars = ax.bar([str(x) for x in yrs], capex, color=["#D9B6B6", "#C77B7B", "#A83E3E", "#800000"])
    for b, v in zip(bars, capex):
        ax.text(b.get_x() + b.get_width()/2, v + 2, f"${v:.1f}B", ha="center", fontsize=9, fontweight="bold")
    ax.set_ylim(0, 132)
    ax.set_ylabel("Cash additions to PP&E ($B)", fontsize=8)
    ax.grid(axis="y", color="#ECE7E4")
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(CHARTS / "msft_capex.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Operating momentum.
    metrics = ["Revenue", "Adj. gross\nmargin", "Adj. op.\nmargin", "Adj. EPS"]
    changes = [83.2, 11.3, 20.1, 320.9]
    fig, ax = plt.subplots(figsize=(7.4, 3.2), dpi=180)
    b = ax.bar(metrics, changes, color=["#800000", "#9E2F2F", "#B45757", "#C87D7D"])
    labs = ["+83%", "+1,130 bps", "+2,010 bps", "+321%"]
    for bar, lab in zip(b, labs):
        ax.text(bar.get_x()+bar.get_width()/2, bar.get_height()+8, lab, ha="center", fontsize=9, fontweight="bold")
    ax.set_ylim(0, 365)
    ax.set_ylabel("FY26 vs FY25 change (%)", fontsize=8)
    ax.grid(axis="y", color="#ECE7E4")
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(CHARTS / "momentum.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Peer cash conversion.
    names = ["LITE", "COHR", "FN"]
    cfo = [751.4, 79.5, 256.7]
    cap = [451.3, 1102.9, 252.5]
    fcf = [300.1, -1023.4, 4.2]
    x = np.arange(3)
    fig, ax = plt.subplots(figsize=(7.5, 3.2), dpi=180)
    w = .25
    ax.bar(x-w, cfo, w, label="CFO", color="#800000")
    ax.bar(x, cap, w, label="Cash capex", color="#C77B7B")
    ax.bar(x+w, fcf, w, label="CFO - capex", color="#C8A951")
    ax.axhline(0, color="#777777", lw=.8)
    ax.set_xticks(x, names)
    ax.set_ylabel("FY26 ($m)", fontsize=8)
    ax.legend(frameon=False, fontsize=8, ncol=3, loc="upper right")
    ax.grid(axis="y", color="#ECE7E4")
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(CHARTS / "peer_fcf.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # Case values.
    cases = ["Bear", "Base", "Bull"]
    values = [331, 1003, 1605]
    colors = ["#B13F3F", "#800000", "#4C1111"]
    fig, ax = plt.subplots(figsize=(7.1, 3.2), dpi=180)
    bars = ax.bar(cases, values, color=colors)
    ax.axhline(893.61, color="#C8A951", ls="--", lw=1.5, label="Sep 17 price $893.61")
    for bar, val in zip(bars, values):
        ax.text(bar.get_x()+bar.get_width()/2, val+35, f"${val:,}", ha="center", fontsize=10, fontweight="bold")
    ax.set_ylim(0, 1780)
    ax.set_ylabel("12-month value / share", fontsize=8)
    ax.legend(frameon=False, fontsize=8)
    ax.grid(axis="y", color="#ECE7E4")
    fig.tight_layout()
    fig.savefig(CHARTS / "case_values.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)

    # DCF sensitivity around workbook base.
    wacc = [9.5, 10.0, 10.5, 11.0, 11.5]
    growth = [2.5, 3.0, 3.5, 4.0, 4.5]
    vals = np.array([
        [561, 526, 495, 468, 443],
        [589, 550, 516, 486, 459],
        [621, 577, 545, 511, 481],
        [657, 608, 566, 530, 497],
        [699, 643, 596, 555, 520],
    ])
    fig, ax = plt.subplots(figsize=(6.8, 3.25), dpi=180)
    im = ax.imshow(vals, cmap="Reds", vmin=430, vmax=910)
    for i in range(vals.shape[0]):
        for j in range(vals.shape[1]):
            ax.text(j, i, f"${vals[i,j]}", ha="center", va="center", fontsize=8,
                    color="white" if vals[i,j] > 610 else "#241F20")
    ax.set_xticks(range(5), [f"{x:.1f}%" for x in wacc])
    ax.set_yticks(range(5), [f"{x:.1f}%" for x in growth])
    ax.set_xlabel("WACC", fontsize=8)
    ax.set_ylabel("Terminal growth", fontsize=8)
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(CHARTS / "dcf_sensitivity.png", bbox_inches="tight", facecolor="white")
    plt.close(fig)


prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(.8)
    return shape


def text(slide, x, y, w, h, value, size=12, color=DARK, bold=False,
         font="Georgia", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=.04, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = value
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    return shape


def rich(slide, x, y, w, h, paras, fill=None, line=None, margin=.08):
    if fill:
        bg = box(slide, x, y, w, h, fill, line or fill)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    for i, item in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item.get("text", "")
        p.level = item.get("level", 0)
        p.font.name = item.get("font", "Georgia")
        p.font.size = Pt(item.get("size", 11))
        p.font.bold = item.get("bold", False)
        p.font.color.rgb = rgb(item.get("color", DARK))
        p.space_after = Pt(item.get("after", 3))
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        if item.get("bullet"):
            p.text = "• " + p.text
    return shape


def title(slide, value, subtitle=None, num=None):
    title_size = 18 if len(value) > 48 else 22
    text(slide, .25, .13, 9.4, .45, value, title_size, MAROON, True)
    box(slide, .25, .62, 9.4, .018, MAROON, MAROON)
    if subtitle:
        text(slide, .3, .68, 9.2, .34, subtitle, 10.5, DARK, False)
    if num is not None:
        text(slide, 9.55, 7.16, .25, .18, str(num), 7.5, MID, False, align=PP_ALIGN.RIGHT)


def footer(slide, source):
    text(slide, .28, 7.13, 9.15, .18, source, 6.6, MID, False, font="Arial")


def section(slide, x, y, w, label):
    box(slide, x, y, w, .31, MAROON, MAROON)
    text(slide, x+.04, y+.015, w-.08, .26, label, 11, WHITE, True, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE)


def add_table(slide, x, y, w, h, data, widths=None, header=True, font_size=9):
    rows, cols = len(data), len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for i, ww in enumerate(widths): table.columns[i].width = Inches(ww)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = cell.margin_right = Inches(.04)
            cell.margin_top = cell.margin_bottom = Inches(.025)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(MAROON if header and r == 0 else (PALE if r % 2 else WHITE))
            cell.border_left = None if hasattr(cell, "border_left") else None
            for p in cell.text_frame.paragraphs:
                p.font.name = "Georgia"; p.font.size = Pt(font_size)
                p.font.bold = header and r == 0
                p.font.color.rgb = rgb(WHITE if header and r == 0 else DARK)
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
    return table


def add_picture(slide, path, x, y, w, h=None):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h) if h else None)


make_charts()

# 1 — Cover
s = prs.slides.add_slide(blank)
box(s, 0, 0, 10, .28, MAROON, MAROON)
box(s, 0, .55, 10, 1.65, MAROON, MAROON)
text(s, .28, .86, 9.44, .55, "Lumentum Holdings (NASDAQ: LITE) | Long Recommendation", 24, WHITE, True,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(s, .35, 1.52, 9.3, .35, "University of Chicago AI-Enabled Investing Competition | 3–12 Month Horizon", 14, WHITE, True,
     align=PP_ALIGN.CENTER)
text(s, 1.15, 3.05, 7.7, .75, "LUMENTUM", 42, MAROON, True, font="Arial", align=PP_ALIGN.CENTER)
text(s, 1.2, 3.80, 7.6, .35, "OPTICS AT THE BOTTLENECK OF AI SCALE", 14, RED, True, font="Arial", align=PP_ALIGN.CENTER)
box(s, .7, 5.54, 8.6, .025, MAROON, MAROON)
text(s, .45, 5.77, 9.1, .42, "Reference Price: $893.61  |  Base 12-Month Value: $1,003  |  Implied Upside: 12.2%", 15, MAROON, True, align=PP_ALIGN.CENTER)
text(s, .7, 6.45, 8.6, .42, "Core view: AI safety commitments can coexist with sustained infrastructure spending; LITE monetizes the optical-content intensity of that spend.", 11.5, DARK, False, align=PP_ALIGN.CENTER)
footer(s, "Reference price: Sep. 17, 2026 close. Model outputs are analyst assumptions; see valuation and risk slides.")

# 2 — Detailed executive summary
s = prs.slides.add_slide(blank); title(s, "Executive Summary", "A narrative shock created a testable 3–12 month earnings-revision setup—not a claim that policy risk is zero.", 2)
section(s, .25, 1.05, 6.22, "ONE THESIS: SAFETY ≠ SPENDING STOP")
rich(s, .25, 1.39, 6.22, 2.48, [
    {"text":"Variant view", "bold":True, "size":11.5, "color":MAROON},
    {"text":"The market overread Dario Amodei’s September pacing proposal as a near-term optical-demand warning. The proposal argues for safeguards and, at the most ambitious level, coordinated pacing; it is not evidence that signed capacity plans or hyperscaler capex were cancelled.", "size":10.4},
    {"text":"Evidence", "bold":True, "size":11.5, "color":MAROON},
    {"text":"We reviewed major safety milestones since 2023—Anthropic’s RSP, Bletchley, Seoul, RSP updates and stricter Claude safeguards. The quantified precedents overlap with rising Microsoft cash PP&E ($28.1B FY23 → $44.5B FY24 → $64.6B FY25 → $115.9B FY26). This rejects inevitability of a pullback; it does not prove zero causal impact.", "size":10.4},
    {"text":"Why Lumentum", "bold":True, "size":11.5, "color":MAROON},
    {"text":"LITE sells the lasers, modules and optical switching that become more important as cluster bandwidth and scale increase. Q4 components grew 103% YoY and systems 123%; 1.6T modules, OCS, CPO lasers, ELS and NPO expand content beyond a single architecture.", "size":10.4},
], fill=CREAM, line=LIGHT_RED)
section(s, 6.70, 1.05, 3.05, "RECOMMENDATION")
rich(s, 6.70, 1.39, 3.05, 2.48, [
    {"text":"LONG LITE", "bold":True, "size":20, "color":MAROON, "align":PP_ALIGN.CENTER},
    {"text":"3–12 months", "bold":True, "size":12, "align":PP_ALIGN.CENTER},
    {"text":"Base value", "size":9, "color":MID, "align":PP_ALIGN.CENTER},
    {"text":"$1,003 / +12.2%", "bold":True, "size":17, "color":MAROON, "align":PP_ALIGN.CENTER},
    {"text":"Bull: $1,605 / +79.6%", "bold":True, "size":10.5, "align":PP_ALIGN.CENTER},
    {"text":"Bear: $331 / −63.0%", "bold":True, "size":10.5, "color":RED, "align":PP_ALIGN.CENTER},
    {"text":"Asymmetric outcomes demand sizing discipline.", "size":8.5, "color":MID, "align":PP_ALIGN.CENTER},
], fill=PALE, line=MAROON)
section(s, .25, 4.03, 9.50, "WHAT MUST HAPPEN / WHAT BREAKS THE PITCH")
add_table(s, .25, 4.39, 9.50, 2.35, [
    ["Earnings path", "Near-term catalysts", "Key risks / invalidation"],
    ["Q1 guide: $1.225–1.275B revenue; 39.5–40.5% adj. op. margin; $4.05–4.35 EPS", "Q1 beat/raise; 1.6T and OCS ramp; purchase-commitment conversion", "Orders weaken while inventory rises; margins miss despite revenue growth"],
    ["Base FY27: $6.10B revenue / $21.83 EPS; FY28: $8.96B / $33.42 EPS", "Consensus revisions toward model; evidence that capex plans remain intact", "Training restrictions become binding; customer concentration / cancellations"],
    ["Base exit: 30× FY28 EPS = $1,003; DCF = $545", "De-risked yields and capacity; mix lifts gross margin", "Valuation compresses; DCF/multiple gap persists; dilution/SBC/capex absorb cash"],
], widths=[3.0, 3.0, 3.5], font_size=8.3)
footer(s, "Sources: Lumentum FY26 results; Amodei essay; Anthropic, Bletchley and Seoul safety materials; Microsoft cash-flow statements; analyst model.")

# 3 — Business overview
s = prs.slides.add_slide(blank); title(s, "Business Overview", "Lumentum is no longer just a telecom component vendor: AI clusters are pulling optics into more links and switching layers.", 3)
section(s, .25, 1.06, 4.58, "REPORTED FY26 / Q4 MOMENTUM")
add_picture(s, CHARTS/"momentum.png", .34, 1.46, 4.35, 2.05)
rich(s, .25, 3.57, 4.58, 1.28, [
    {"text":"FY26 revenue: $3.014B (+83% YoY)", "bold":True, "size":10.5, "color":MAROON},
    {"text":"Q4 revenue: $1.006B; adjusted GM 50.4%; adjusted operating margin 36.6%.", "size":9.7},
    {"text":"Q1 FY27 midpoint reaches $1.25B with ~40% adjusted operating margin—more than a quarter ahead of management’s target-model timing.", "size":9.7},
], fill=CREAM, line=LIGHT_RED)
section(s, 5.08, 1.06, 4.67, "WHERE THE OPTICAL CONTENT SITS")
add_table(s, 5.08, 1.46, 4.67, 2.12, [
    ["Platform", "Q4 rev.", "AI relevance"],
    ["Components", "$649.4m", "High-power CPO lasers; external laser sources; NPO"],
    ["Systems", "$356.9m", "1.6T cloud modules; optical circuit switching"],
    ["Total", "$1,006.3m", "+109% YoY; both platforms expanding"],
], widths=[1.05, .85, 2.77], font_size=8.6)
rich(s, 5.08, 3.72, 4.67, 1.13, [
    {"text":"Economic logic", "bold":True, "size":10.5, "color":MAROON},
    {"text":"More accelerators require more interconnect bandwidth; higher speeds shorten the reach of copper and increase optical links. LITE can win from both pluggable upgrades and emerging in-rack optical architectures.", "size":9.7},
], fill=CREAM, line=LIGHT_RED)
section(s, .25, 5.05, 9.50, "BALANCE SHEET + CASH CONVERSION")
add_table(s, .25, 5.42, 9.50, 1.21, [
    ["Cash + short-term investments", "Debt carrying value", "FY26 CFO", "FY26 cash capex", "CFO − capex"],
    ["$2.738B", "$1.637B", "$751m", "$451m", "$300m"],
], widths=[1.95, 1.85, 1.75, 1.85, 2.10], font_size=10)
footer(s, "Source: Lumentum FY26 earnings release and FY26 10-K. Non-GAAP metrics exclude items detailed by the company.")

# 4 — Thesis
s = prs.slides.add_slide(blank); title(s, "Thesis: Safety Commitments Can Coexist with Capex", "The investable claim is narrower than “safety does not matter”: announcements have not mechanically produced spending retrenchment.", 4)
section(s, .25, 1.05, 4.55, "WHAT AMODEI PROPOSED")
rich(s, .25, 1.42, 4.55, 2.16, [
    {"text":"A ladder—not an immediate unilateral shutdown", "bold":True, "size":11, "color":MAROON},
    {"text":"1. Ban narrow dangerous uses.", "size":10.2, "bullet":True},
    {"text":"2. Test frontier models for cyber, biology and alignment risks.", "size":10.2, "bullet":True},
    {"text":"3. Explore a verifiable “speed limit” on recursive self-improvement.", "size":10.2, "bullet":True},
    {"text":"4. A broad pacing/pause is described as difficult and unlikely soon without high-confidence verification.", "size":10.2, "bullet":True},
], fill=CREAM, line=LIGHT_RED)
section(s, 5.02, 1.05, 4.73, "WHY SPENDING CAN CONTINUE")
rich(s, 5.02, 1.42, 4.73, 2.16, [
    {"text":"Safety spend and compute spend are complements", "bold":True, "size":11, "color":MAROON},
    {"text":"Evaluations, inference, interpretability, redundancy and cyber controls still require compute and networking.", "size":10.2, "bullet":True},
    {"text":"Hyperscalers invest for cloud customers, enterprise workloads and inference—not only frontier training.", "size":10.2, "bullet":True},
    {"text":"Committed campuses, power, chips and networking have long lead times; policy headlines do not instantly cancel the installed plan.", "size":10.2, "bullet":True},
], fill=CREAM, line=LIGHT_RED)
section(s, .25, 3.83, 9.50, "OBSERVED MARKET REACTION")
add_picture(s, CHARTS/"price_reaction.png", .35, 4.17, 5.52, 2.13)
rich(s, 6.03, 4.19, 3.72, 2.10, [
    {"text":"Tape read", "bold":True, "size":11, "color":MAROON},
    {"text":"LITE fell 9.9% on Sep. 14, then recovered to $919 on Sep. 16 and closed Sep. 17 at $893.61.", "size":10},
    {"text":"Interpretation", "bold":True, "size":11, "color":MAROON},
    {"text":"The reaction was sharp but did not coincide with a public Lumentum guide cut or a hyperscaler capex cancellation. We treat this as a narrative stress test, not proof of causality.", "size":10},
], fill=PALE, line=MAROON)
footer(s, "Sources: Dario Amodei, “We Must Pace the Frontier” (Sep. 2026); StockAnalysis LITE price history accessed Sep. 18, 2026.")

# 5 — Precedents
s = prs.slides.add_slide(blank); title(s, "Precedent Review: Commitments Rose Alongside Infrastructure Spend", "We screened recurring safety announcements; three milestones permit a clean, quantified public-capex cross-check.", 5)
section(s, .25, 1.05, 5.35, "EVENT STUDY (DESCRIPTIVE, NOT CAUSAL)")
add_table(s, .25, 1.42, 5.35, 3.00, [
    ["Safety milestone", "Commitment", "Observed spending context"],
    ["Sep. 2023 Anthropic RSP", "Pause training/deployment if safeguards fail", "AWS later added $4B to Anthropic partnership; financing, not capex"],
    ["Nov. 2023 Bletchley", "International risk coordination", "MSFT cash PP&E FY23 $28.1B → FY24 $44.5B (+58%)"],
    ["May 2024 Seoul", "No development/deployment in extreme cases without mitigation", "MSFT cash PP&E FY24 $44.5B → FY25 $64.6B (+45%)"],
    ["2024–26 updates", "RSP revisions, ASL-3 safeguards, risk reporting", "MSFT FY26 cash PP&E $115.9B; FY27 capex expected to grow"],
], widths=[1.20, 1.75, 2.40], font_size=7.7)
section(s, 5.83, 1.05, 3.92, "MICROSOFT CASH PP&E")
add_picture(s, CHARTS/"msft_capex.png", 5.91, 1.48, 3.73, 2.15)
rich(s, 5.83, 3.70, 3.92, .72, [
    {"text":"The sequence falsifies “safety announcement ⇒ automatic capex decline.” It cannot isolate what capex would have been absent the announcement.", "size":9, "color":MAROON, "bold":True},
], fill=PALE, line=MAROON)
section(s, .25, 4.65, 9.50, "BROADER SCREEN: SIX HIGH-PROFILE SAFETY MOMENTS, NO DISCLOSED OPTICAL-CAPEX CANCELLATION")
add_table(s, .25, 5.02, 9.50, 1.54, [
    ["2023 RSP", "Bletchley", "2024 RSP reflection", "Seoul commitments", "2024 RSP update", "2025 ASL-3 safeguards"],
    ["Framework launched", "28 countries coordinate", "Red-line plans refined", "Frontier commitments", "Thresholds/safeguards", "Claude Opus 4 controls"],
    ["Compute partnership expanded", "Capex rose", "Capex rose", "Capex rose", "No public optical cut", "Stricter release + continued scaling"],
], widths=[1.58]*6, font_size=7.7)
footer(s, "Sources: Anthropic policy archive; UK Bletchley/Seoul declarations; Microsoft FY23–FY26 cash flows. Windows overlap events; no causal coefficient claimed.")

# 6 — Why Lumentum
s = prs.slides.add_slide(blank); title(s, "Why Lumentum Is Best Positioned for This Thesis", "The preferred exposure combines direct optical leverage, accelerating margins, positive cash conversion and multiple architecture paths.", 6)
section(s, .25, 1.05, 4.65, "LITE VS. OPTICAL / MANUFACTURING PEERS")
add_picture(s, CHARTS/"peer_fcf.png", .33, 1.42, 4.45, 2.10)
add_table(s, .25, 3.60, 4.65, 1.36, [
    ["FY26", "LITE", "COHR", "FN"],
    ["Revenue", "$3.0B", "$7.1B", "$4.6B"],
    ["CFO − capex", "+$300m", "−$1,023m", "+$4m"],
    ["Net cash / (debt)", "+$1.10B", "−$1.24B", "+$0.88B"],
], widths=[1.55, 1.03, 1.03, 1.04], font_size=8.7)
section(s, 5.13, 1.05, 4.62, "FOUR REASONS LITE FITS THE VIEW")
rich(s, 5.13, 1.42, 4.62, 3.54, [
    {"text":"1 | Pure optical-content torque", "bold":True, "size":11, "color":MAROON},
    {"text":"Components and systems directly monetize bandwidth, reach and switching complexity; FN is a contract manufacturer and COHR carries broader industrial exposure.", "size":9.6},
    {"text":"2 | Multiple architecture paths", "bold":True, "size":11, "color":MAROON},
    {"text":"1.6T pluggables, OCS, CPO lasers, ELS and NPO reduce dependence on one interconnect design winning.", "size":9.6},
    {"text":"3 | Operating leverage is already visible", "bold":True, "size":11, "color":MAROON},
    {"text":"Q4 adjusted gross margin reached 50.4% and operating margin 36.6%; the thesis does not start from hoped-for profitability.", "size":9.6},
    {"text":"4 | Strategic demand signal", "bold":True, "size":11, "color":MAROON},
    {"text":"NVIDIA’s nonexclusive multi-year purchase commitment and $2B investment support capacity access—but do not guarantee quarterly revenue.", "size":9.6},
], fill=CREAM, line=LIGHT_RED)
section(s, .25, 5.18, 9.50, "LIMIT OF THE CLAIM")
text(s, .38, 5.58, 9.20, .75, "“Best positioned” means the cleanest risk/reward expression among the reviewed public comparables—not that Lumentum is operationally safest. Customer concentration, ramp yields, cancellations, dilution and a premium valuation remain material.", 11, MAROON, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
footer(s, "Sources: FY26 company filings/releases; NVIDIA–Lumentum partnership announcement; consensus snapshot dated Aug. 31, 2026.")

# 7 — Forecast
s = prs.slides.add_slide(blank); title(s, "Earnings Bridge: The 3–12 Month Debate Is About Revisions", "Base case assumes guidance execution, continued optical-content growth and margin expansion—not a new hyperscaler capex step-up.", 7)
section(s, .25, 1.05, 5.48, "QUARTERLY BASE CASE")
add_table(s, .25, 1.43, 5.48, 2.56, [
    ["$m except EPS", "Q1:27", "Q2:27", "Q3:27", "Q4:27", "FY27"],
    ["Revenue", "1,250", "1,420", "1,620", "1,810", "6,101"],
    ["Seq. growth", "24.2%", "13.6%", "14.1%", "11.7%", "102.4% YoY"],
    ["Adj. gross margin", "52.5%", "54.0%", "55.0%", "55.5%", "54.4%"],
    ["Adj. op. margin", "40.0%", "42.2%", "43.8%", "44.7%", "42.9%"],
    ["Adj. EPS", "$4.17", "$5.00", "$5.92", "$6.75", "$21.83"],
], widths=[1.38, .82, .82, .82, .82, .82], font_size=8.3)
rich(s, .25, 4.18, 5.48, 1.17, [
    {"text":"Near-term model discipline", "bold":True, "size":10.5, "color":MAROON},
    {"text":"Q1 is calibrated to management guidance. Later quarters are analyst assumptions—not disclosed backlog. Base assumes no price/mix benefit and no incremental training-order haircut.", "size":9.7},
], fill=CREAM, line=LIGHT_RED)
section(s, 5.98, 1.05, 3.77, "BASE VS. CONSENSUS")
add_table(s, 5.98, 1.43, 3.77, 1.70, [
    ["FY27", "Model", "Aug. consensus", "Delta"],
    ["Revenue", "$6.10B", "$6.32B", "−3.5%"],
    ["Adj. EPS", "$21.83", "$21.67", "+0.8%"],
    ["Implication", "Margin-led", "Street", "Slight EPS beat"],
], widths=[1.05, .91, 1.10, .71], font_size=8.3)
section(s, 5.98, 3.40, 3.77, "FY28 BASE STEP-UP")
rich(s, 5.98, 3.78, 3.77, 1.57, [
    {"text":"Revenue $8.96B | EPS $33.42", "bold":True, "size":13, "color":MAROON, "align":PP_ALIGN.CENTER},
    {"text":"Components +45%; systems +50%; 56% adjusted GM; fixed opex +12%.", "size":9.6, "align":PP_ALIGN.CENTER},
    {"text":"This is the key underwriting risk: FY28 visibility must build within the holding period.", "bold":True, "size":9.4, "align":PP_ALIGN.CENTER},
], fill=PALE, line=MAROON)
section(s, .25, 5.62, 9.50, "THE CHECKPOINTS")
text(s, .35, 5.98, 9.28, .62, "Q1: hit guide and hold ~40% op. margin  |  Q2–Q3: prove 1.6T/OCS ramps and gross-margin leverage  |  6–12 months: convert commitments to shipments without inventory or capex overruns", 10.5, DARK, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
footer(s, "Source: Lumentum guidance; analyst model in Lumentum_Investment_Model.xlsx. Product volumes, training shares and price/mix are scenario assumptions.")

# 8 — Valuation
s = prs.slides.add_slide(blank); title(s, "Valuation: A Multiple-Driven 12-Month Long, with a DCF Warning", "The base target is not a DCF output; the disagreement is the principal valuation risk and is shown rather than reconciled away.", 8)
section(s, .25, 1.05, 5.05, "12-MONTH CASES")
add_picture(s, CHARTS/"case_values.png", .35, 1.42, 4.75, 2.18)
add_table(s, .25, 3.70, 5.05, 1.36, [
    ["Case", "FY28 EPS", "Exit P/E", "Value", "Return"],
    ["Bear", "$13.22", "25×", "$331", "−63.0%"],
    ["Base", "$33.42", "30×", "$1,003", "+12.2%"],
    ["Bull", "$45.86", "35×", "$1,605", "+79.6%"],
], widths=[1.05, 1.05, .90, 1.05, 1.00], font_size=8.5)
section(s, 5.53, 1.05, 4.22, "DCF CROSS-CHECK")
rich(s, 5.53, 1.42, 4.22, 1.15, [
    {"text":"Base DCF: $545 / share", "bold":True, "size":16, "color":MAROON, "align":PP_ALIGN.CENTER},
    {"text":"−39% vs. reference price | 72% of EV from terminal value", "size":9.5, "align":PP_ALIGN.CENTER},
    {"text":"10.5% WACC | 3.5% g | 20% terminal ROIC", "size":9, "color":MID, "align":PP_ALIGN.CENTER},
], fill=PALE, line=MAROON)
add_picture(s, CHARTS/"dcf_sensitivity.png", 5.65, 2.74, 3.98, 2.25)
section(s, .25, 5.30, 9.50, "VALUATION INTERPRETATION")
rich(s, .25, 5.67, 9.50, .86, [
    {"text":"The 3–12 month long works if earnings visibility moves faster than the multiple compresses. At 30×, the stock needs FY28 EPS of $35.74 for a 20% return—7% above the base model. The DCF says today’s price already capitalizes aggressive long-duration growth; therefore position size and quarterly checkpoints matter more than a single “fair value.”", "size":10.4, "bold":True, "color":MAROON, "align":PP_ALIGN.CENTER},
], fill=CREAM, line=LIGHT_RED)
footer(s, "Source: analyst model. Scenario probabilities are not shown because outcomes are highly non-linear; reference price is Sep. 17, 2026.")

# 9 — Catalysts and risks
s = prs.slides.add_slide(blank); title(s, "Catalysts, Risks and Falsification", "A short holding period requires observable tests; every catalyst has a paired failure signal.", 9)
section(s, .25, 1.05, 4.60, "CATALYSTS (3–12 MONTHS)")
rich(s, .25, 1.42, 4.60, 4.42, [
    {"text":"Next earnings: guide execution", "bold":True, "size":11, "color":MAROON},
    {"text":"Revenue at/above $1.25B midpoint, ~40% adjusted operating margin and EPS toward the high end establish the earnings floor.", "size":10},
    {"text":"1.6T cloud-module + OCS ramp", "bold":True, "size":11, "color":MAROON},
    {"text":"Mix and scale can lift gross margin faster than consensus even if headline capex remains flat.", "size":10},
    {"text":"CPO / ELS / NPO milestones", "bold":True, "size":11, "color":MAROON},
    {"text":"Additional qualifications or orders expand optical content per accelerator cluster.", "size":10},
    {"text":"No hyperscaler capex retrenchment", "bold":True, "size":11, "color":MAROON},
    {"text":"Reaffirmed spending converts the September narrative debate into an earnings debate.", "size":10},
    {"text":"Consensus revision cycle", "bold":True, "size":11, "color":MAROON},
    {"text":"FY28 estimates moving toward the model’s $33.42 EPS supports the 30× base exit.", "size":10},
], fill=CREAM, line=LIGHT_RED)
section(s, 5.12, 1.05, 4.63, "RISKS / INVALIDATION")
rich(s, 5.12, 1.42, 4.63, 4.42, [
    {"text":"Policy becomes binding", "bold":True, "size":11, "color":RED},
    {"text":"A coordinated, verifiable training cap—or unilateral customer cuts—can reduce exposed demand; the historical precedents are weaker analogues.", "size":10},
    {"text":"Execution and concentration", "bold":True, "size":11, "color":RED},
    {"text":"Customer cancellations, yield constraints, delayed qualifications or an architecture shift can overwhelm sector spending.", "size":10},
    {"text":"Cash conversion / dilution", "bold":True, "size":11, "color":RED},
    {"text":"High capex, working capital and recurring SBC can make non-GAAP EPS overstate economic value.", "size":10},
    {"text":"Valuation", "bold":True, "size":11, "color":RED},
    {"text":"At ~41× FY27 consensus EPS, LITE has little room for a demand or margin miss; base DCF is well below market.", "size":10},
    {"text":"Sell / reassess if", "bold":True, "size":11, "color":RED},
    {"text":"Shipments miss while inventory rises; gross margin weakens despite scale; hyperscalers cut networking plans; or FY28 EPS visibility stays below ~$30.", "size":10},
], fill=PALE, line=RED)
section(s, .25, 6.05, 9.50, "POSITIONING DISCIPLINE")
text(s, .38, 6.38, 9.20, .36, "Build only after the next guide confirms demand; size for the bear-case drawdown; do not use the DCF as support for the long.", 11.2, MAROON, True, align=PP_ALIGN.CENTER)
footer(s, "Source: company risk factors, management guidance and analyst scenario analysis.")

# 10 — Assumptions
s = prs.slides.add_slide(blank); title(s, "Model Assumptions", "Blue-cell workbook inputs are scenario judgments; reported facts and dated consensus snapshots are kept separate.", 10)
section(s, .25, 1.05, 9.50, "KEY OPERATING AND VALUATION ASSUMPTIONS")
add_table(s, .25, 1.42, 9.50, 3.36, [
    ["Driver", "Bear", "Base", "Bull", "Rationale / limitation"],
    ["FY27 revenue / EPS", "$4.66B / $13.55", "$6.10B / $21.83", "$7.05B / $26.68", "Q1 calibrated to guidance; later quarters are analyst assumptions"],
    ["FY28 component growth", "+5%", "+45%", "+60%", "No disclosed units or ASPs; volume-index approach"],
    ["FY28 systems growth", "+8%", "+50%", "+65%", "Captures 1.6T / OCS ramp scenarios"],
    ["FY28 adj. gross margin", "46%", "56%", "59%", "Scale/mix; benchmark against Q4 FY26 50.4%"],
    ["FY27 capex / revenue", "18%", "15%", "16%", "Economic FCF remains capex intensive"],
    ["Inventory days FY27", "150", "115", "105", "Execution risk concentrated in ramp and conversion"],
    ["Training-order haircut", "10–20% after Q1", "0%", "0%", "Applied only to assumed exposed shares: components 40%, systems 35%"],
    ["Exit P/E", "25×", "30×", "35×", "Sensitivity assumption; not peer-derived certainty"],
], widths=[1.55, 1.12, 1.12, 1.12, 4.59], font_size=7.8)
section(s, .25, 5.03, 4.62, "DCF ASSUMPTIONS")
add_table(s, .25, 5.40, 4.62, 1.28, [
    ["WACC", "Terminal g", "Terminal ROIC", "Shares"],
    ["10.5%", "3.5%", "20%", "102m"],
    ["Sensitivity", "2.5–4.5%", "Reinvestment explicit", "Proxy"],
], widths=[1.14, 1.14, 1.20, 1.14], font_size=8.5)
section(s, 5.13, 5.03, 4.62, "ACCOUNTING TREATMENT")
rich(s, 5.13, 5.40, 4.62, 1.28, [
    {"text":"SBC deducted as an economic cost; no double-counted future dilution. PP&E uses a 7-year remaining life and half-year capex convention. Operating AP is normalized because reported AP includes capex financing.", "size":9.1, "align":PP_ALIGN.CENTER},
], fill=CREAM, line=LIGHT_RED)
footer(s, "See Lumentum_Investment_Model.xlsx: Assumptions, Quarterly, DCF, Actuals and Evidence tabs.")

# 11 — DCF detail
s = prs.slides.add_slide(blank); title(s, "Appendix: DCF Build", "Economic free cash flow deducts recurring SBC and models capex, depreciation and working capital explicitly.", 11)
section(s, .25, 1.05, 9.50, "BASE CASE — USD MILLIONS EXCEPT PER SHARE")
add_table(s, .25, 1.42, 9.50, 3.20, [
    ["", "FY27", "FY28", "FY29", "FY30", "FY31", "FY32", "FY33"],
    ["Revenue", "6,101", "8,961", "11,821", "14,417", "16,753", "18,565", "20,050"],
    ["Adj. operating margin", "42.9%", "45.8%", "46.5%", "45.8%", "45.0%", "44.2%", "43.2%"],
    ["Economic op. profit", "2,437", "3,838", "5,138", "6,172", "7,043", "7,643", "8,069"],
    ["Cash capex", "915", "1,254", "1,418", "1,442", "1,340", "1,300", "1,203"],
    ["Increase in NWC", "447", "568", "545", "468", "385", "252", "298"],
    ["Economic UFCF", "940", "1,716", "2,691", "3,655", "4,646", "5,382", "5,836"],
    ["PV of UFCF", "850", "1,405", "1,995", "2,451", "2,820", "2,957", "2,901"],
], widths=[1.90]+[1.085]*7, font_size=8.0)
section(s, .25, 4.88, 4.60, "VALUATION BRIDGE")
add_table(s, .25, 5.25, 4.60, 1.36, [
    ["PV explicit", "PV terminal", "EV", "Net cash", "Equity / sh."],
    ["$15,380", "$39,146", "$54,526", "$1,101", "$545"],
], widths=[.92]*5, font_size=8.7)
section(s, 5.12, 4.88, 4.63, "TERMINAL ECONOMICS")
rich(s, 5.12, 5.25, 4.63, 1.36, [
    {"text":"Terminal FCF = NOPAT × (1 − g / ROIC)", "bold":True, "size":11, "color":MAROON, "align":PP_ALIGN.CENTER},
    {"text":"This forces growth to consume reinvestment. Terminal value is 71.8% of EV, so the DCF is highly sensitive and should be treated as a risk flag—not a precise target.", "size":9.1, "align":PP_ALIGN.CENTER},
], fill=CREAM, line=LIGHT_RED)
footer(s, "Source: analyst model, DCF tab. Cash/debt use FY26 reported balances; forecast is not a full GAAP three-statement model.")

# 12 — Sources / AI disclosure
s = prs.slides.add_slide(blank); title(s, "Appendix: Sources, Method and AI-Use Disclosure", "Every major claim is traceable; unresolved data are identified rather than silently estimated.", 12)
section(s, .25, 1.05, 4.65, "PRIMARY / COMPANY SOURCES")
rich(s, .25, 1.42, 4.65, 3.76, [
    {"text":"Lumentum", "bold":True, "size":11, "color":MAROON},
    {"text":"FY26 earnings release; FY26 10-K; Q1 FY27 guidance; NVIDIA strategic partnership announcement.", "size":9.3},
    {"text":"Peers", "bold":True, "size":11, "color":MAROON},
    {"text":"Coherent FY26 release/presentation; Fabrinet FY26 10-K and earnings materials.", "size":9.3},
    {"text":"Safety / policy", "bold":True, "size":11, "color":MAROON},
    {"text":"Dario Amodei, “We Must Pace the Frontier”; Anthropic RSP archive; Bletchley Declaration; Seoul Frontier AI Safety Commitments.", "size":9.3},
    {"text":"Capex / market data", "bold":True, "size":11, "color":MAROON},
    {"text":"Microsoft FY23–FY26 cash-flow statements and FY27 outlook; dated StockAnalysis price/consensus snapshots.", "size":9.3},
], fill=CREAM, line=LIGHT_RED)
section(s, 5.13, 1.05, 4.62, "METHOD / LIMITATIONS")
rich(s, 5.13, 1.42, 4.62, 3.76, [
    {"text":"What the evidence supports", "bold":True, "size":11, "color":MAROON},
    {"text":"Prior safety commitments did not mechanically force disclosed infrastructure spending lower; LITE has direct optical exposure and strong reported operating momentum.", "size":9.3},
    {"text":"What it does not support", "bold":True, "size":11, "color":RED},
    {"text":"No causal estimate of safety policy on capex; no claim that all datacenter revenue is AI; no disclosed training/inference split; no guaranteed quarterly revenue from purchase commitments.", "size":9.3},
    {"text":"AI-use disclosure", "bold":True, "size":11, "color":MAROON},
    {"text":"AI assisted source retrieval, spreadsheet checks, scenario calculations, writing and slide production. Forecasts are analyst-authored assumptions. No management interviews, proprietary channel checks or claimed human verification were performed.", "size":9.3},
], fill=PALE, line=MAROON)
section(s, .25, 5.44, 9.50, "FILES DELIVERED")
text(s, .35, 5.83, 9.28, .75, "Final_Lumentum_Bull_Thesis.pptx  |  Lumentum_Investment_Model.xlsx  |  Source archive and model remain editable in the project folder", 12, MAROON, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
footer(s, "Prepared Sep. 18, 2026. Investment presentation for educational competition use; not investment advice.")


dest = OUT / "Final_Lumentum_Bull_Thesis.pptx"
prs.save(dest)
print(dest)
print(f"slides={len(prs.slides)}")
