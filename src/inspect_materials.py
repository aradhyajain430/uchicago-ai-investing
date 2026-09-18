"""Dump Excel, existing pitch, and template PDFs for the final deck."""
from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation
import fitz

ROOT = Path(r"C:\Users\Dell\Desktop\uchicagoaicomp")
OUT = ROOT / "research" / "extracts"
OUT.mkdir(parents=True, exist_ok=True)

xlsx = Path(r"C:\Users\Dell\Downloads\Lumentum_Investment_Model.xlsx")
wb = load_workbook(xlsx, data_only=True)
wb_f = load_workbook(xlsx, data_only=False)

lines = [f"sheets: {wb.sheetnames}"]
for name in wb.sheetnames:
    ws = wb[name]
    ws_f = wb_f[name]
    lines.append(f"\n===== SHEET {name} dims={ws.dimensions} =====")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 120), max_col=min(ws.max_column or 1, 18)):
        vals = []
        for c in row:
            if c.value is not None:
                vals.append(f"{c.coordinate}={c.value!r}")
        if vals:
            lines.append(" | ".join(vals))
(OUT / "excel_values.txt").write_text("\n".join(lines), encoding="utf-8")

formula_lines = []
for name in wb_f.sheetnames:
    ws = wb_f[name]
    formula_lines.append(f"\n===== FORMULAS {name} =====")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 120), max_col=min(ws.max_column or 1, 18)):
        vals = []
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                vals.append(f"{c.coordinate}={c.value}")
            elif c.value is not None and not isinstance(c.value, str):
                pass
        if vals:
            formula_lines.append(" | ".join(vals))
(OUT / "excel_formulas.txt").write_text("\n".join(formula_lines), encoding="utf-8")

pptx_path = ROOT / "slideshow" / "Round1_Pitch.pptx"
prs = Presentation(str(pptx_path))
p_lines = [f"slides={len(prs.slides)} size={prs.slide_width}x{prs.slide_height}"]
for i, slide in enumerate(prs.slides, 1):
    p_lines.append(f"\n===== SLIDE {i} =====")
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if text.strip():
                p_lines.append(text)
                p_lines.append("---")
(OUT / "round1_pitch.txt").write_text("\n".join(p_lines), encoding="utf-8")

for pdf_name in ["1-UPenn-Merge (1).pdf", "LCV Capital-OTIS-ExecSum.pdf"]:
    pdf = ROOT / "slideshow" / pdf_name
    doc = fitz.open(pdf)
    pdf_lines = [f"pages={doc.page_count} file={pdf_name}"]
    for i, page in enumerate(doc, 1):
        pdf_lines.append(f"\n===== PAGE {i} =====")
        pdf_lines.append(page.get_text("text"))
    safe = pdf_name.replace(" ", "_").replace("(", "").replace(")", "")
    (OUT / f"{safe}.txt").write_text("\n".join(pdf_lines), encoding="utf-8")
    print(pdf_name, "pages", doc.page_count)

print("excel sheets", wb.sheetnames)
print("pptx slides", len(prs.slides))
